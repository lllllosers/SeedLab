"""Validate the fixed legacy workbook; apply only to an empty, temporary test DB.

Never opens the configured application Data Root. Source workbooks are read only.
The CLI requires an explicitly selected, migrated test DB and an existing owner.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
from zoneinfo import ZoneInfo

BACKEND = Path(__file__).resolve().parents[2] / "backend"
sys.path.insert(0, str(BACKEND))

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.session import make_engine
from app.models import (Experiment, ExperimentMaterial, ExperimentProtocol, GerminationDish,
                        GerminationObservation, MeasurementTimepoint, SeedlingMeasurement,
                        SeedlingSample, SeedLot, Taxon, User)
from app.services.common import next_code, record
from app.services.experiment_identity import experiment_batch_month, next_experiment_code
from app.services.local_time import local_datetime, utc_naive
from app.services.ordering import dish_display_number, sample_display_number
from app.services.workbook_export import build

SOURCE_SHA256 = "1099d3074570ef83f53c03ed24ffb413848dfd256969258d24fb41d049255f3e"
REVISION = "d2e7a46b910c"
DAYS = (3, 7, 14)
START = date(2026, 8, 3)
SCIENTIFIC_NAMES = {197: "Hypochaeris radicata", 198: "Bromus catharticus",
                    199: "Cynoglossum amabile", 200: "Himalaiella deltoidea"}
SOWN_DATES = {number: datetime(2026, 8, 3) for number in (51, 65, 72, 126, 143)}
EXPECTED = {"materials": 200, "taxa": 198, "samples": 1665,
            "dag": {3: 1655, 7: 1655, 14: 1645}, "measurements": 4955,
            "slots": 4995, "missing_subtraction": 40, "missing_enumeration": 40,
            "missing_per_dag": {3: 10, 7: 10, 14: 20}, "root_zeros": 6,
            "shoot_zeros": 283, "one_sided": 0, "negative": 0, "max_delay": 43}
BUSINESS_TABLES = ("taxa", "seed_lots", "experiments", "experiment_protocols",
                   "experiment_materials", "germination_dishes", "germination_observations",
                   "seedling_samples", "seedling_measurements", "measurement_timepoints", "import_jobs")


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


@dataclass(frozen=True)
class LegacyMaterial:
    number: int
    common_name: str
    scientific_name: str
    source: str | None
    source_code: str | None
    sown_at: datetime


@dataclass(frozen=True)
class LegacySample:
    material_number: int
    number: int
    germinated_at: datetime
    lengths: tuple[tuple[Decimal | None, Decimal | None], ...]


@dataclass(frozen=True)
class LegacyData:
    materials: tuple[LegacyMaterial, ...]
    samples: tuple[LegacySample, ...]
    metrics: dict


def source_utc(value: datetime) -> datetime:
    require(isinstance(value, datetime), "原表日期格式不正确，请核对源文件")
    if value.tzinfo is None:
        value = value.replace(tzinfo=ZoneInfo(get_settings().seedlab_timezone))
    return value.astimezone(timezone.utc).replace(tzinfo=None)


def read_source(path: Path) -> LegacyData:
    require(path.is_absolute(), "请提供原始 Excel 的绝对路径")
    with path.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
        require(digest == SOURCE_SHA256, f"原始 Excel SHA256 不一致：{digest}")
        stream.seek(0)
        book = load_workbook(stream, read_only=True, data_only=True)
        try:
            metadata = list(book["种子信息"].values)
            length_rows = list(book["根-苗长统计表"].values)
            settings = {row[0]: row[1] for row in list(book["试验设置"].values)[1:]}
        finally:
            book.close()
    require(metadata[0][:4] == ("编号", "物种", "拉丁文名", "来源"), "原表材料列发生变化，请停止核账")
    require(length_rows[0][:6] == ("物种编号", "物种名称", "种子编号", "样本ID(勿动）", "置床时间", "发芽时间"),
            "原表根苗长列发生变化，请停止核账")
    require(length_rows[1][6:12] == ("3DAG", "7DAG", "14DAG", "3DAG", "7DAG", "14DAG"),
            "原表 DAG 列发生变化，请停止核账")
    require(settings.get("总种子数") == 50 and settings.get("根苗长取样数") == 10
            and settings.get("默认重复") == "R1" and str(settings.get("根苗长测定节点")) == "3,7,14",
            "原表试验设置发生变化，请停止核账")
    info = {}
    for row in metadata[1:]:
        if row[0] is None:
            continue
        number = int(row[0])
        require(number not in info, f"历史材料 {number:03d} 重复")
        info[number] = row
    require(set(info) == set(range(1, 201)), "历史材料编号必须完整且唯一，为 001—200")
    grouped = defaultdict(list)
    for row in length_rows[2:]:
        if row[0] is not None:
            grouped[int(row[0])].append(row)
    require(set(grouped) == set(info), "材料信息与根苗长表的材料编号不一致")
    materials, samples = [], []
    one_sided = negative = 0
    for number, row in sorted(info.items()):
        raw_name = row[2]
        require((raw_name is None) == (number in SCIENTIFIC_NAMES), f"材料 {number:03d} 缺失学名与人工映射不一致")
        scientific = SCIENTIFIC_NAMES[number] if raw_name is None else " ".join(str(raw_name).replace("_", " ").split())
        records = grouped[number]
        require(len(records) == 10 and {int(r[2]) for r in records} == set(range(1, 11)),
                f"材料 {number:03d} 的幼苗序号不是完整的 1—10")
        dates = {r[4] for r in records if r[4] is not None}
        require((not dates) == (number in SOWN_DATES) and len(dates) <= 1,
                f"材料 {number:03d} 置床日期与人工确认范围不一致")
        sown = source_utc(SOWN_DATES[number] if not dates else next(iter(dates)))
        materials.append(LegacyMaterial(number, str(row[1]), scientific, row[3],
                                        str(row[6]) if row[6] is not None else None, sown))
        for r in records:
            lengths = tuple(tuple(Decimal(str(v)) if v is not None else None for v in (r[6+i], r[9+i]))
                            for i in range(3))
            for root, shoot in lengths:
                one_sided += (root is None) != (shoot is None)
                negative += sum(v is not None and v < 0 for v in (root, shoot))
                require(all(v is None or v == v.quantize(Decimal(".01")) for v in (root, shoot)),
                        f"材料 {number:03d} 的长度超出现有两位小数精度，不得自动舍入")
            if r[5] is None:
                require(all(pair == (None, None) for pair in lengths), f"材料 {number:03d} 存在无发芽时间的测定")
                continue
            germinated = source_utc(r[5])
            require(germinated >= sown, f"材料 {number:03d} 发芽时间早于置床时间")
            samples.append(LegacySample(number, int(r[2]), germinated, lengths))
    dag = {day: sum(sample.lengths[i][0] is not None and sample.lengths[i][1] is not None
                    for sample in samples) for i, day in enumerate(DAYS)}
    # Independent path: enumerate each missing sample-DAG pair, not a reused subtraction.
    missing = [(sample.material_number, sample.number, day) for sample in samples
               for i, day in enumerate(DAYS) if sample.lengths[i] == (None, None)]
    missing_per_dag = Counter(key[2] for key in missing)
    sown_by_number = {m.number: m.sown_at for m in materials}
    metrics = {"materials": len(materials), "taxa": len({m.scientific_name for m in materials}),
               "samples": len(samples), "dag": dag, "measurements": sum(dag.values()),
               "slots": len(samples)*len(DAYS), "missing_subtraction": len(samples)*len(DAYS)-sum(dag.values()),
               "missing_enumeration": len(missing), "missing_per_dag": {day: missing_per_dag[day] for day in DAYS},
               "root_zeros": sum(pair[0] == 0 for s in samples for pair in s.lengths),
               "shoot_zeros": sum(pair[1] == 0 for s in samples for pair in s.lengths),
               "one_sided": one_sided, "negative": negative,
               "max_delay": max((s.germinated_at - sown_by_number[s.material_number]).days for s in samples)}
    differences = {key: {"expected": expected, "actual": metrics[key]}
                   for key, expected in EXPECTED.items() if metrics[key] != expected}
    require(not differences, "源数据基线不一致：" + json.dumps(differences, ensure_ascii=False))
    shared = defaultdict(list)
    for material in materials:
        shared[material.scientific_name].append(material.number)
    require({name: numbers for name, numbers in shared.items() if len(numbers) > 1}
            == {"Lepidium apetalum": [33, 153], "Lappula myosotis": [55, 56]}, "共享物种与确认清单不一致")
    require(local_datetime(min(sown_by_number.values())).date() == START, "最早置床日期不一致")
    metrics["all_dag_empty_samples"] = [[s.material_number, s.number] for s in samples
                                         if all(pair == (None, None) for pair in s.lengths)]
    obtained_counts = Counter(s.material_number for s in samples)
    metrics["obtained_per_material"] = {f"{m.number:03d}": obtained_counts[m.number] for m in materials}
    require(metrics["all_dag_empty_samples"] == [[27, n] for n in range(1, 11)], "全 DAG 未测样本清单不一致")
    return LegacyData(tuple(materials), tuple(samples), metrics)


def temporary_database(path: Path) -> Path:
    path = path.resolve()
    temporary_root = Path(tempfile.gettempdir()).resolve()
    require(path.is_relative_to(temporary_root) and path.suffix.lower() == ".db",
            "仅允许系统临时目录内的一次性 .db 测试数据库；正式数据目录不可使用")
    require(not any(part.casefold() == "seedlabdata" for part in path.parts), "禁止访问正式 SeedLabData")
    require(not any((parent / name).exists() for parent in path.parents
                    for name in ("seedlab.json", "installation.json", "config/seedlab.json", "config/installation.json")),
            "目标属于运行中的数据目录，请改用一次性临时测试库")
    require(path.is_file(), "请先在系统临时目录准备迁移到当前版本且含负责人的空业务测试数据库")
    wal = Path(str(path) + "-wal")
    require(not wal.exists() or wal.stat().st_size == 0, "请先关闭临时测试库连接并完成写入，再运行核验")
    return path


def inspect_target(path: Path, owner: str) -> str:
    # Immutable read-only mode cannot create WAL/SHM files or change the DB.
    with closing(sqlite3.connect(path.as_uri() + "?mode=ro&immutable=1", uri=True)) as db:
        require(db.execute("SELECT version_num FROM alembic_version").fetchall() == [(REVISION,)],
                "测试数据库迁移版本不一致，请先更新临时测试库")
        occupied = {table: db.execute(f"SELECT count(*) FROM {table}").fetchone()[0] for table in BUSINESS_TABLES}
        require(not any(occupied.values()), "测试数据库已有业务数据，请另建空业务临时库；工具不会清空或覆盖")
        owners = db.execute("SELECT id FROM users WHERE (username=? OR id=?) AND is_active=1", (owner, owner)).fetchall()
        require(len(owners) == 1, "请指定临时测试库中唯一且可用的实验负责人")
        return owners[0][0]


def source_measurements(data: LegacyData) -> dict:
    return {(s.material_number, s.number, day): (s.germinated_at, s.germinated_at + timedelta(days=day), root, shoot)
            for s in data.samples for day, (root, shoot) in zip(DAYS, s.lengths) if root is not None and shoot is not None}


def reconcile_database(db: Session, data: LegacyData, experiment: Experiment) -> dict:
    expected_counts = {"taxa": 198, "seed_lots": 200, "experiments": 1, "experiment_materials": 200,
                       "germination_dishes": 200, "germination_observations": 0,
                       "seedling_samples": 1665, "seedling_measurements": 4955, "measurement_timepoints": 3}
    counts = {table: db.connection().exec_driver_sql(f"SELECT count(*) FROM {table}").scalar()
              for table in expected_counts}
    require(counts == expected_counts, f"临时库数量核账不一致：{counts}")
    require(experiment.code == "GER-202608-001" and experiment.experiment_type == "GER"
            and experiment.status == "active" and experiment.planned_start_date == START,
            "实验编号、类型、状态或计划开始日期不一致")
    materials = list(db.scalars(select(ExperimentMaterial)))
    dishes = list(db.scalars(select(GerminationDish)))
    numbers = {m.id: m.experiment_number for m in materials}
    material_by_number = {m.experiment_number: m for m in materials}
    dish_numbers = {d.id: numbers[d.material_id] for d in dishes}
    actual_sown = {dish_numbers[d.id]: utc_naive(d.sown_at) for d in dishes}
    require(actual_sown == {m.number: m.sown_at for m in data.materials}, "实际置床时间逐值核账不一致")
    lots = {lot.id: lot for lot in db.scalars(select(SeedLot))}
    taxa = {taxon.id: taxon for taxon in db.scalars(select(Taxon))}
    for source in data.materials:
        material = material_by_number[source.number]
        lot = lots[material.seed_lot_id]
        require(material.label == source.common_name and lot.source == source.source
                and lot.source_code == source.source_code and taxa[lot.taxon_id].scientific_name == source.scientific_name,
                f"材料 {source.number:03d} 基础信息核账不一致")
    samples = list(db.scalars(select(SeedlingSample)))
    sample_keys = {s.id: (dish_numbers[s.dish_id], s.sample_number) for s in samples}
    actual_samples = {sample_keys[s.id]: utc_naive(s.germinated_at) for s in samples}
    require(actual_samples == {(s.material_number, s.number): s.germinated_at for s in data.samples},
            "幼苗编号或实际发芽时间逐值核账不一致")
    require(all(s.source_observation_id is None for s in samples), "历史幼苗不得关联虚构巡检")
    timepoints = {t.id: t.day_after_germination for t in db.scalars(select(MeasurementTimepoint))}
    require(sorted(timepoints.values()) == list(DAYS), "测定时间点不一致")
    measurements = list(db.scalars(select(SeedlingMeasurement)))
    actual = {(*sample_keys[m.sample_id], timepoints[m.timepoint_id]):
              (actual_samples[sample_keys[m.sample_id]], utc_naive(m.measured_at), m.root_length_mm, m.shoot_length_mm)
              for m in measurements}
    require(actual == source_measurements(data), "根苗长、DAG 或实际测定时间逐值核账不一致")
    require(all(not m.root_unavailable and not m.shoot_unavailable for m in measurements), "历史长度不得补 NA")
    missing_enumerated = [(key, day) for key in actual_samples for day in DAYS if (*key, day) not in actual]
    missing_subtraction = len(actual_samples)*len(timepoints)-len(measurements)
    require(missing_subtraction == len(missing_enumerated) == 40, "临时库两种缺失核账必须均为 40")
    require(db.connection().exec_driver_sql("PRAGMA foreign_key_check").all() == [], "临时库关联核验失败")
    return {"counts": counts, "value_differences": 0, "missing_subtraction": missing_subtraction,
            "missing_enumeration": len(missing_enumerated),
            "dag": dict(sorted(Counter(key[2] for key in actual).items())),
            "missing_per_dag": dict(sorted(Counter(day for _, day in missing_enumerated).items())),
            "root_zeros": sum(m.root_length_mm == 0 for m in measurements),
            "shoot_zeros": sum(m.shoot_length_mm == 0 for m in measurements)}


def reconcile_workbook(stream: BytesIO, data: LegacyData) -> dict:
    book = load_workbook(stream, read_only=True, data_only=True)
    try:
        long_rows = list(book["04_幼苗测定长表"].values)[1:]
        wide_rows = list(book["05_幼苗测定宽表"].values)
        require(wide_rows[0][8:14] == tuple(column for day in DAYS for column in (f"RL{day}", f"SL{day}"))
                and wide_rows[0][14:] == ("发芽判定时间", "是否已有实际幼苗", "计划取样序号", "取样范围", "计划培养皿重复"),
                "导出宽表 DAG 列不一致")
        source_samples = {(s.material_number, s.number): s for s in data.samples}
        source_values = source_measurements(data)
        planned_samples = {(m.number, number) for m in data.materials for number in range(1, 11)}
        planned_stages = {(*key, day) for key in planned_samples for day in DAYS}
        expected_long = {}
        expected_status = {}
        for key in planned_stages:
            source_sample = source_samples.get(key[:2])
            expected_long[key] = source_values.get(key, (
                source_sample.germinated_at if source_sample else None, None, None, None))
            expected_status[key] = "已测定" if key in source_values else "无测定记录" if source_sample else "无实际幼苗"
        actual_long = {}
        actual_status = {}
        for row in long_rows:
            number, sample = int(row[2]), int(row[4].rsplit("-", 1)[1])
            key = (number, sample, row[9])
            require(row[3] == dish_display_number(number, 1, 1)
                    and row[4] == sample_display_number(number, 1, 1, sample), "导出长表现场编号不一致")
            require(key not in actual_long, "导出长表存在重复测定")
            require(all(row[i] is None or (isinstance(row[i], (int, float)) and not isinstance(row[i], bool)) for i in (13, 15)),
                    "导出长表长度必须为数值或空白，真实 0 不能变成文字")
            actual_long[key] = (utc_naive(datetime.fromisoformat(row[8])) if row[8] else None,
                                utc_naive(datetime.fromisoformat(row[11])) if row[11] else None,
                                Decimal(str(row[13])) if row[13] is not None else None,
                                Decimal(str(row[15])) if row[15] is not None else None)
            actual_status[key] = row[18]
            require(row[19] == sample and row[20] == "每皿" and row[21] == 1, "导出长表计划取样结构不一致")
            if row[18] == "已测定":
                require(row[12] == 0 and row[14] == row[16] == "已测", "历史测定日期或长度状态导出不一致")
            else:
                require(row[11:14] == (None, None, None) and row[15] is None
                        and row[14] == row[16] == row[18], "导出缺失槽位不得补时间、长度或 NA")
            expected_planned = (local_datetime(source_samples[key[:2]].germinated_at).date()
                                + timedelta(days=key[2])).isoformat() if key[:2] in source_samples else None
            require(row[10] == expected_planned, "导出长表计划日期与实际幼苗匹配不一致")
        require(actual_long == expected_long and actual_status == expected_status, "导出长表计划槽位或逐值核账不一致")
        statuses = Counter(actual_status.values())
        require(statuses == {"已测定": 4955, "无测定记录": 40, "无实际幼苗": 1005}
                and len(actual_long) == 6000, "4955 / 40 / 1005 / 6000 导出闭环不一致")
        actual_wide = {}
        actual_sample_facts = {}
        for row in wide_rows[1:]:
            number, sample = int(row[2]), int(row[4].rsplit("-", 1)[1])
            key = (number, sample)
            require(key not in actual_wide, "导出宽表存在重复幼苗")
            require(row[3] == dish_display_number(number, 1, 1)
                    and row[4] == sample_display_number(number, 1, 1, sample), "导出宽表现场编号不一致")
            require(all(v is None or (isinstance(v, (int, float)) and not isinstance(v, bool)) for v in row[8:14]),
                    "导出宽表长度必须为数值或空白，不能变成文字或 NA")
            actual_wide[key] = tuple(Decimal(str(v)) if v is not None else None for v in row[8:14])
            actual_sample_facts[key] = (utc_naive(datetime.fromisoformat(row[14])) if row[14] else None, row[15])
            require(row[16] == sample and row[17] == "每皿" and row[18] == 1, "导出宽表计划取样结构不一致")
        expected_wide = {key: tuple(v for pair in source_samples[key].lengths for v in pair)
                         if key in source_samples else (None,)*6 for key in planned_samples}
        expected_sample_facts = {key: (source_samples[key].germinated_at, "是") if key in source_samples else (None, "否")
                                 for key in planned_samples}
        require(actual_wide == expected_wide and actual_sample_facts == expected_sample_facts,
                "导出宽表计划槽位、实际幼苗或逐值核账不一致")
        require(Counter(key[0] for key in actual_wide) == {m.number: 10 for m in data.materials},
                "导出宽表必须包含全部 200 份材料，每份材料 10 个计划槽位")
        require(all(actual_wide[(27, n)] == (None,)*6 and actual_sample_facts[(27, n)][1] == "是"
                    for n in range(1, 11)), "材料 027 的实际幼苗不能误标为无实际幼苗")
        rate_rows = list(book["02_发芽率汇总"].values)[1:]
        require(len(rate_rows) == 200 and all(row[9:11] == (None, None) for row in rate_rows),
                "无巡检时累计发芽和发芽率必须导出为空白")
        require(book["03_发芽原始记录"].max_row == 1, "历史实验不得导出虚构巡检")
        require(all("-M" not in str(value) for sheet in book for row in sheet.values for value in row),
                "导出中出现培养皿内部编号")
        return {"long_rows": len(actual_long), "wide_rows": len(actual_wide), "value_differences": 0,
                "planned_sample_slots": len(planned_samples), "obtained_sample_slots": len(source_samples),
                "absent_sample_slots": len(planned_samples)-len(source_samples),
                "planned_measurement_slots": len(planned_stages), "measured_slots": statuses["已测定"],
                "unmeasured_actual_slots": statuses["无测定记录"], "absent_sample_measurement_slots": statuses["无实际幼苗"],
                "all_dag_empty_rows": sum(all(v is None for v in values) for values in actual_wide.values()),
                "all_dag_empty_actual_rows": sum(all(v is None for v in actual_wide[key]) for key in source_samples),
                "root_zeros": sum(values[i] == 0 for values in actual_wide.values() for i in (0, 2, 4)),
                "shoot_zeros": sum(values[i] == 0 for values in actual_wide.values() for i in (1, 3, 5))}
    finally:
        book.close()


def apply(data: LegacyData, path: Path, owner_id: str) -> dict:
    engine = make_engine(f"sqlite:///{path.as_posix()}")
    try:
        with engine.connect() as connection:
            connection.exec_driver_sql("BEGIN IMMEDIATE")
            try:
                with Session(bind=connection) as db:
                    require(all(connection.exec_driver_sql(f"SELECT count(*) FROM {table}").scalar() == 0
                                for table in BUSINESS_TABLES), "临时库已有业务数据，请重新准备空业务库")
                    owner = db.get(User, owner_id)
                    require(owner is not None and owner.is_active, "临时库负责人不可用")
                    earliest = min(m.sown_at for m in data.materials)
                    experiment = Experiment(code=next_experiment_code(db, "GER", experiment_batch_month(START)),
                        experiment_type="GER", name="200份材料历史种子萌发试验（导入能力验证）",
                        description="仅用于临时测试；不含历史发芽巡检、发芽率或干重。", owner_id=owner_id,
                        planned_start_date=START, status="active", started_at=earliest, numbering_locked_at=earliest)
                    db.add(experiment)
                    db.flush()
                    db.add(ExperimentProtocol(experiment_id=experiment.id, seeds_per_dish=50, replicate_count=1,
                        sample_count=10, sample_scope="per_dish", sampling_rule="first_germinated",
                        summary="原表每份材料置床50粒、取样10株、单重复；未补充原表缺失的观察周期和发芽判定标准。"))
                    timepoints = {}
                    for day in DAYS:
                        point = MeasurementTimepoint(experiment_id=experiment.id, day_after_germination=day)
                        db.add(point)
                        db.flush()
                        timepoints[day] = point.id
                    taxa, dishes = {}, {}
                    for source in data.materials:
                        taxon = taxa.get(source.scientific_name)
                        if taxon is None:
                            taxon = Taxon(code=next_code(db, Taxon, "SP-"), scientific_name=source.scientific_name,
                                          common_name=source.common_name)
                            db.add(taxon)
                            db.flush()
                            taxa[source.scientific_name] = taxon
                        lot = SeedLot(code=next_code(db, SeedLot, f"LOT-{local_datetime(earliest).year}-", 3),
                                      taxon_id=taxon.id, source=source.source, source_code=source.source_code,
                                      notes=f"历史材料 {source.number:03d}：{source.common_name}；仅用于导入能力验证。")
                        db.add(lot)
                        db.flush()
                        material = ExperimentMaterial(experiment_id=experiment.id, seed_lot_id=lot.id,
                            experiment_number=source.number, display_order=source.number-1, label=source.common_name)
                        db.add(material)
                        db.flush()
                        dish = GerminationDish(material_id=material.id,
                            code=f"{experiment.code}-M{source.number:03d}-R01", replicate_no=1,
                            label="R1", seed_count=50, sown_at=source.sown_at)
                        db.add(dish)
                        db.flush()
                        dishes[source.number] = dish.id
                    for source in data.samples:
                        sample = SeedlingSample(dish_id=dishes[source.material_number], sample_number=source.number,
                                                germinated_at=source.germinated_at)
                        db.add(sample)
                        db.flush()
                        for day, (root, shoot) in zip(DAYS, source.lengths):
                            if root is not None and shoot is not None:
                                db.add(SeedlingMeasurement(sample_id=sample.id, timepoint_id=timepoints[day],
                                    root_length_mm=root, shoot_length_mm=shoot,
                                    measured_at=source.germinated_at+timedelta(days=day)))
                    record(db, owner_id, "import", "Experiment", experiment.id, None,
                           {"code": experiment.code, "source_sha256": SOURCE_SHA256, "test_only": True})
                    db.flush()
                    database_report = reconcile_database(db, data, experiment)
                    workbook_report = reconcile_workbook(build(db, [experiment.id]), data)
                    result = {"experiment_code": experiment.code, "database": database_report, "workbook": workbook_report}
                connection.commit()
                return result
            except Exception:
                connection.rollback()
                raise
    finally:
        engine.dispose()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="200份历史材料的导入能力验证；仅使用一次性临时测试库")
    parser.add_argument("--source", type=Path, required=True, help="原始 Excel 绝对路径（只读）")
    target = parser.add_mutually_exclusive_group(required=True)
    target.add_argument("--database", type=Path, help="系统临时目录内已迁移的空业务 .db 测试库")
    target.add_argument("--data-root", type=Path, help="系统临时目录内的一次性测试目录，数据库为 seedlab.db")
    parser.add_argument("--owner", required=True, help="临时库中已有负责人的用户名或身份")
    parser.add_argument("--apply", action="store_true", help="明确写入临时测试库；默认只读核验")
    args = parser.parse_args(argv)
    try:
        data = read_source(args.source)
        path = temporary_database(args.database if args.database else args.data_root / "seedlab.db")
        owner_id = inspect_target(path, args.owner)
        result = {"status": "PASS", "mode": "apply" if args.apply else "dry-run",
                  "source_sha256": SOURCE_SHA256, "source": data.metrics}
        if args.apply:
            result.update(apply(data, path, owner_id))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ValueError, OSError) as exc:
        print(f"核验停止：{exc}", file=sys.stderr)
    except (sqlite3.Error, SQLAlchemyError):
        print("核验停止：临时测试库无法读取或写入。请检查迁移版本、负责人和空业务状态；未完成的写入已回滚。", file=sys.stderr)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
