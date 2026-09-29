import csv
import io
from datetime import datetime
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.api.schemas import AuditOut, PasswordReset, UserCreate, UserOut, UserPatch
from app.core.auth import admin_user, current_user, hash_password, require_password
from app.db.session import get_db
from app.models import AuditLog, Experiment, ImportJob, SeedLot, SessionToken, Taxon, User
from app.services.common import commit_or_conflict, flush_or_conflict, next_code, record, require_entity


router = APIRouter(tags=["system"])


def safe_user_state(user: User) -> dict:
    return {"username": user.username, "display_name": user.display_name,
            "is_admin": user.is_admin, "is_active": user.is_active}


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    count = lambda model: db.scalar(select(func.count()).select_from(model)) or 0
    return {
        "taxa": count(Taxon), "seed_lots": count(SeedLot), "experiments": count(Experiment),
        "active_experiments": db.scalar(select(func.count()).select_from(Experiment).where(Experiment.status == "active")) or 0,
        "recent_experiments": [{"id": x.id, "code": x.code, "name": x.name, "status": x.status} for x in db.scalars(select(Experiment).order_by(Experiment.created_at.desc()).limit(5))],
        "recent_actions": [{"id": x.id, "action": x.action, "entity_type": x.entity_type, "created_at": x.created_at.isoformat()} for x in db.scalars(select(AuditLog).order_by(AuditLog.created_at.desc()).limit(5))],
    }


@router.get("/audit-logs", response_model=list[AuditOut])
def audit_logs(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    return db.scalars(select(AuditLog).order_by(AuditLog.created_at.desc()).limit(100)).all()


@router.get("/users", response_model=list[UserOut])
def users(db: Session = Depends(get_db), _admin: User = Depends(admin_user)):
    return db.scalars(select(User).order_by(User.created_at)).all()


@router.post("/users", response_model=UserOut, status_code=201)
def create_user(data: UserCreate, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    require_password(data.password)
    user = User(username=data.username, display_name=data.display_name.strip(),
                password_hash=hash_password(data.password), is_admin=data.is_admin,
                must_change_password=True)
    if not user.display_name:
        raise HTTPException(422, "请填写显示姓名")
    db.add(user)
    flush_or_conflict(db)
    record(db, admin.id, "create", "User", user.id, None, safe_user_state(user))
    commit_or_conflict(db)
    return user


@router.patch("/users/{user_id}", response_model=UserOut)
def update_user(user_id: str, data: UserPatch, db: Session = Depends(get_db),
                admin: User = Depends(admin_user)):
    target = require_entity(db, User, user_id)
    patch = data.model_dump(exclude_unset=True)
    if any(value is None for value in patch.values()):
        raise HTTPException(422, "用户姓名、角色和状态不能为空")
    if "display_name" in patch:
        patch["display_name"] = patch["display_name"].strip()
        if not patch["display_name"]:
            raise HTTPException(422, "请填写显示姓名")
    if target.id == admin.id and patch.get("is_active") is False:
        raise HTTPException(409, "不能停用自己的账号")
    if target.id == admin.id and patch.get("is_admin") is False:
        raise HTTPException(409, "不能取消自己的管理员权限")
    losing_admin = target.is_admin and target.is_active and (
        patch.get("is_admin", target.is_admin) is False or
        patch.get("is_active", target.is_active) is False)
    if losing_admin and (db.scalar(select(func.count()).select_from(User).where(
            User.is_admin.is_(True), User.is_active.is_(True))) or 0) <= 1:
        raise HTTPException(409, "不能移除最后一个启用的管理员")
    before = safe_user_state(target)
    for key, value in patch.items():
        setattr(target, key, value)
    invalidated_sessions = patch.get("is_active") is False
    if invalidated_sessions:
        db.execute(delete(SessionToken).where(SessionToken.user_id == target.id))
    after = safe_user_state(target)
    if after != before:
        record(db, admin.id, "update", "User", target.id, before, after)
    if after != before or invalidated_sessions:
        commit_or_conflict(db)
    return target


@router.post("/users/{user_id}/reset-password", response_model=UserOut)
def reset_user_password(user_id: str, data: PasswordReset,
                        db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    target = require_entity(db, User, user_id)
    if target.id == admin.id:
        raise HTTPException(409, "请通过个人改密入口修改自己的密码")
    require_password(data.password)
    before = {"initial_change_required": target.must_change_password}
    target.password_hash = hash_password(data.password)
    target.must_change_password = True
    db.execute(delete(SessionToken).where(SessionToken.user_id == target.id))
    record(db, admin.id, "update", "User", target.id, before,
           {"initial_change_required": True, "credential_event": "admin_reset"})
    commit_or_conflict(db)
    return target


@router.get("/export/taxa.csv")
def export_taxa(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["code", "scientific_name", "common_name", "family", "is_active", "notes"])
    for item in db.scalars(select(Taxon).order_by(Taxon.code)):
        writer.writerow([item.code, item.scientific_name, item.common_name or "", item.family or "", item.is_active, item.notes or ""])
    return StreamingResponse(iter(["\ufeff" + output.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": 'attachment; filename="seedlab-taxa.csv"'})


@router.get("/export/taxa-template.xlsx")
def taxa_template(_user: User = Depends(current_user)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "物种导入"
    sheet.append(("scientific_name", "common_name", "family"))
    for column, width in {"A": 32, "B": 24, "C": 24}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="seedlab-taxa-template.xlsx"'})


@router.post("/import/taxa", status_code=201)
async def import_taxa(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(current_user)):
    filename = file.filename or ""
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(422, "请上传 .xlsx 文件")
    content = await file.read(2_000_001)
    if len(content) > 2_000_000:
        raise HTTPException(413, "文件不能超过 2 MB")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = [str(value or "").strip() for value in next(rows)]
        if "scientific_name" not in header:
            raise ValueError("首行缺少学名列，请按导入格式填写")
        entries = []
        for line, values in enumerate(rows, start=2):
            row = dict(zip(header, values))
            name = str(row.get("scientific_name") or "").strip()
            if not name:
                if any(value is not None for value in values):
                    raise ValueError(f"第 {line} 行缺少学名")
                continue
            entries.append({"scientific_name": name, "common_name": str(row.get("common_name") or "").strip() or None,
                            "family": str(row.get("family") or "").strip() or None})
        workbook.close()
    except (OSError, BadZipFile) as exc:
        raise HTTPException(422, "无法读取物种文件，请确认上传的是有效的 Excel 文件") from exc
    except (ValueError, StopIteration) as exc:
        raise HTTPException(422, f"文件格式错误：{exc}") from exc
    if len(entries) > 500:
        raise HTTPException(422, "一次最多导入 500 行")
    names = [entry["scientific_name"] for entry in entries]
    if len(set(names)) != len(names) or db.scalar(select(Taxon.id).where(Taxon.scientific_name.in_(names)).limit(1)):
        raise HTTPException(409, "学名在文件中重复，或系统中已有同名物种；请检查后重新导入")
    job = ImportJob(user_id=user.id, filename=filename, status="completed", total_rows=len(entries), successful_rows=len(entries))
    db.add(job)
    start = int(next_code(db, Taxon, "SP-").removeprefix("SP-"))
    for offset, entry in enumerate(entries):
        item = Taxon(code=f"SP-{start + offset:04d}", **entry)
        db.add(item)
        flush_or_conflict(db)
        record(db, user.id, "import", "Taxon", item.id, None, {"code": item.code, **entry})
    commit_or_conflict(db)
    return {"id": job.id, "imported": len(entries)}


SEED_LOT_COLUMNS = ("taxon_code", "common_name", "scientific_name", "source", "quantity", "notes")


@router.get("/export/seed-lots-template.xlsx")
def seed_lots_template(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "种子批次导入"
    sheet.append(SEED_LOT_COLUMNS)
    for taxon in db.scalars(select(Taxon).where(Taxon.is_active.is_(True)).order_by(Taxon.code)):
        sheet.append((taxon.code, taxon.common_name, taxon.scientific_name, None, None, None))
        for cell in sheet[sheet.max_row][:3]:
            cell.data_type = "s"
    for column, width in {"A": 18, "B": 24, "C": 32, "D": 30, "E": 16, "F": 42}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="seedlab-seed-lots-template.xlsx"'})


@router.post("/import/seed-lots", status_code=201)
async def import_seed_lots(file: UploadFile = File(...), db: Session = Depends(get_db),
                           user: User = Depends(current_user)):
    filename = file.filename or ""
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(422, "请选择 Excel 格式的种子批次文件（.xlsx）")
    content = await file.read(5_000_001)
    if len(content) > 5_000_000:
        raise HTTPException(413, "文件不能超过 5 MB")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            header = [str(value or "").strip() for value in next(rows)]
            if "taxon_code" not in header or len(header) != len(set(header)):
                raise ValueError("首行需要包含物种编号列，且列名不能重复；请使用下载的模板")
            entries = []
            for line, values in enumerate(rows, start=2):
                if all(value is None or str(value).strip() == "" for value in values):
                    continue
                row = dict(zip(header, values))
                entries.append((line, row))
                if len(entries) > 500:
                    raise ValueError("一次最多导入 500 行，请分批上传")
        finally:
            workbook.close()
    except (OSError, BadZipFile) as exc:
        raise HTTPException(422, "无法读取种子批次文件，请确认上传的是有效的 Excel 文件") from exc
    except (ValueError, StopIteration) as exc:
        raise HTTPException(422, f"无法读取种子批次文件：{exc}") from exc
    if not entries:
        raise HTTPException(422, "文件中没有需要导入的种子批次")

    codes = {str(row.get("taxon_code") or "").strip() for _, row in entries}
    taxa = {item.code: item for item in db.scalars(select(Taxon).where(Taxon.code.in_(codes)))}
    validated = []
    errors = []
    for line, row in entries:
        code = str(row.get("taxon_code") or "").strip()
        taxon = taxa.get(code)
        if not code:
            errors.append(f"第 {line} 行：请填写物种编号")
        elif taxon is None:
            errors.append(f"第 {line} 行：找不到物种编号 {code}")
        elif not taxon.is_active:
            errors.append(f"第 {line} 行：物种 {code} 已停用，不能添加种子批次")

        source = str(row.get("source") or "").strip() or None
        if source and len(source) > 255:
            errors.append(f"第 {line} 行：来源不能超过 255 个字符")
        raw_quantity = row.get("quantity")
        quantity = None
        if raw_quantity is not None and str(raw_quantity).strip() != "":
            try:
                if isinstance(raw_quantity, bool) or int(raw_quantity) < 0 or int(raw_quantity) > 2**63 - 1 or (
                    isinstance(raw_quantity, (int, float)) and raw_quantity != int(raw_quantity)
                ) or (isinstance(raw_quantity, str) and raw_quantity.strip() != str(int(raw_quantity))):
                    raise ValueError
                quantity = int(raw_quantity)
            except (TypeError, ValueError, OverflowError):
                errors.append(f"第 {line} 行：数量请填写大于或等于 0 的整数（最多 19 位），或留空")
        validated.append((taxon, source, quantity, str(row.get("notes") or "").strip() or None))
    if errors:
        raise HTTPException(422, "整表未导入。请修正以下行后重新上传：\n" + "\n".join(errors))

    prefix = f"LOT-{datetime.now().year}-"
    start = int(next_code(db, SeedLot, prefix, 3).rsplit("-", 1)[1])
    job = ImportJob(user_id=user.id, filename=filename, status="completed",
                    total_rows=len(validated), successful_rows=len(validated))
    db.add(job)
    for offset, (taxon, source, quantity, notes) in enumerate(validated):
        item = SeedLot(code=f"{prefix}{start + offset:03d}", taxon_id=taxon.id,
                       source=source, quantity=quantity, notes=notes)
        db.add(item)
        flush_or_conflict(db)
        record(db, user.id, "import", "SeedLot", item.id, None,
               {"code": item.code, "taxon_code": taxon.code, "source": source,
                "quantity": quantity, "notes": notes})
    commit_or_conflict(db)
    return {"id": job.id, "imported": len(validated)}
