"""Single and multi-experiment workbook built from the same source facts."""

from datetime import datetime, timezone
from collections import defaultdict
from io import BytesIO

from app.contracts.errors import NotFoundError, ValidationError
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (Experiment, ExperimentMaterial, GerminationDish,
                        GerminationObservation, SeedLot, Taxon)
from app.services.ordering import display_number, field_number, material_key
from app.services.measurement_slots import build_measurement_dataset
from app.services.local_time import local_date, local_datetime
from app.version import VERSION


def _date(value):
    return value.isoformat() if value is not None else None


def build(db: Session, experiment_ids: list[str]) -> BytesIO:
    if not experiment_ids or len(experiment_ids) != len(set(experiment_ids)):
        raise ValidationError("请至少选择一个不重复的实验")
    experiments = {item.id: item for item in db.scalars(select(Experiment).where(Experiment.id.in_(experiment_ids)))}
    if len(experiments) != len(experiment_ids):
        raise NotFoundError("部分实验不存在，请刷新列表后重试")
    rows = db.execute(select(ExperimentMaterial, SeedLot, Taxon)
                      .join(SeedLot, ExperimentMaterial.seed_lot_id == SeedLot.id)
                      .join(Taxon, SeedLot.taxon_id == Taxon.id)
                      .where(ExperimentMaterial.experiment_id.in_(experiment_ids))).all()
    rows = sorted(rows, key=lambda row: (*material_key(row[2], row[1]), experiments[row[0].experiment_id].code,
                                          row[0].id))
    if not rows:
        raise ValidationError("所选实验还没有材料，无法生成工作簿")
    canonical = build_measurement_dataset(db, experiment_ids)
    slots_by_material = defaultdict(list)
    stages_by_slot = defaultdict(list)
    for slot in canonical.seedling_slots:
        slots_by_material[slot.material_id].append(slot)
    for stage in canonical.rows:
        stages_by_slot[stage.slot.key].append(stage)
    workbook = Workbook()
    materials_sheet = workbook.active
    materials_sheet.title = "01_材料总表"
    rate_sheet = workbook.create_sheet("02_发芽率汇总")
    observation_sheet = workbook.create_sheet("03_发芽原始记录")
    long_sheet = workbook.create_sheet("04_幼苗测定长表")
    wide_sheet = workbook.create_sheet("05_幼苗测定宽表")
    explanation = workbook.create_sheet("06_导出说明")
    materials_sheet.append(("汇总编号", "来源实验", "原实验编号", "中文名", "学名", "科", "属", "生活型",
                            "原始材料编号", "系统物种编号", "系统种子批次编号", "来源", "采集/获得日期", "数量", "备注"))
    rate_sheet.append(("汇总编号", "来源实验", "原实验编号", "中文名", "学名", "原始材料编号",
                       "种子批次", "实际已置床培养皿数", "实际置床种子数", "累计发芽数", "发芽率（%）"))
    observation_sheet.append(("汇总编号", "来源实验", "原实验编号", "培养皿现场编号",
                              "中文名", "学名", "实际置床时间", "巡检时间", "本次新增发芽数", "累计发芽数", "发芽率"))
    long_sheet.append(("汇总编号", "来源实验", "原实验编号", "培养皿现场编号", "幼苗编号", "位置标签",
                       "中文名", "学名", "发芽判定时间", "DAG", "计划测定日期", "实际测定时间",
                       "延迟天数", "根长（mm）", "根长状态", "苗长（mm）", "苗长状态", "备注",
                       "数据状态", "计划取样序号", "取样范围", "计划培养皿重复"))
    all_dag = sorted({point.day_after_germination for point in canonical.stages})
    wide_sheet.append(("汇总编号", "来源实验", "原实验编号", "培养皿现场编号", "幼苗编号", "位置标签", "中文名", "学名") +
                      tuple(column for day in all_dag for column in (f"RL{day}", f"SL{day}")) +
                      ("发芽判定时间", "是否已有实际幼苗", "计划取样序号", "取样范围", "计划培养皿重复"))
    material_ids = [item.id for item, _, _ in rows]
    dishes = list(db.scalars(select(GerminationDish).where(GerminationDish.material_id.in_(material_ids))))
    dishes_by_material = {}
    for dish in dishes:
        dishes_by_material.setdefault(dish.material_id, []).append(dish)
    observations = list(db.scalars(select(GerminationObservation).where(
        GerminationObservation.dish_id.in_([dish.id for dish in dishes])))) if dishes else []
    observations_by_dish = {}
    for observation in observations:
        observations_by_dish.setdefault(observation.dish_id, []).append(observation)
    for index, (material, lot, taxon) in enumerate(rows, start=1):
        summary_number = display_number(index)
        experiment = experiments[material.experiment_id]
        source_name = f"{experiment.name}（{experiment.code}）"
        original_number = display_number(material.experiment_number) if material.experiment_number else None
        materials_sheet.append((summary_number, source_name, original_number, taxon.common_name,
                                taxon.scientific_name, taxon.family, taxon.genus, taxon.life_form,
                                lot.source_code, taxon.code, lot.code, lot.source, _date(lot.collected_at),
                                lot.quantity, lot.notes))
        relevant = sorted(dishes_by_material.get(material.id, []), key=lambda dish: dish.replicate_no)
        sown = [dish for dish in relevant if dish.sown_at is not None and dish.cancelled_at is None]
        actual_seeds = sum(dish.seed_count for dish in sown)
        recorded = [observation for dish in sown for observation in observations_by_dish.get(dish.id, [])]
        germinated = sum(observation.new_germinated_count for observation in recorded) if recorded else None
        rate_sheet.append((summary_number, source_name, original_number, taxon.common_name,
                           taxon.scientific_name, lot.source_code, lot.code, len(sown), actual_seeds,
                           germinated, round(germinated / actual_seeds * 100, 2)
                           if germinated is not None and actual_seeds else None))
        for dish in relevant:
            replicate_count = max(item.replicate_no for item in relevant)
            dish_number = field_number(material, replicate_count, dish.replicate_no)
            count = 0
            for observation in sorted(observations_by_dish.get(dish.id, []), key=lambda item: (item.observed_at, item.id)):
                count += observation.new_germinated_count
                observation_sheet.append((summary_number, source_name, original_number, dish_number,
                                          taxon.common_name, taxon.scientific_name,
                                          _date(dish.sown_at), _date(observation.observed_at),
                                          observation.new_germinated_count, count,
                                          round(count / dish.seed_count * 100, 2)))
        for slot in slots_by_material[material.id]:
            germinated_at = slot.germinated_at
            germinated_text = local_datetime(germinated_at).isoformat() if germinated_at else None
            position = slot.position_label
            scope_text = "每皿" if slot.sample_scope == "per_dish" else "每材料"
            values = {}
            for stage in stages_by_slot[slot.key]:
                measurement = stage if stage.measurement_exists else None
                day = stage.stage.day_after_germination
                root = float(measurement.root_length_mm) if measurement and measurement.root_length_mm is not None else None
                shoot = float(measurement.shoot_length_mm) if measurement and measurement.shoot_length_mm is not None else None
                planned = stage.scheduled_date
                long_sheet.append((summary_number, source_name, original_number, slot.dish_number,
                    slot.seedling_number, position, taxon.common_name, taxon.scientific_name,
                    germinated_text, day, _date(planned),
                    local_datetime(measurement.measured_at).isoformat() if measurement else None,
                    (local_date(measurement.measured_at) - planned).days if measurement and planned else None,
                    root, ("无法测量" if measurement.root_unavailable else "已测") if measurement else stage.data_status,
                    shoot, ("无法测量" if measurement.shoot_unavailable else "已测") if measurement else stage.data_status,
                    measurement.notes if measurement else None, stage.data_status, slot.planned_number, scope_text, slot.replicate_no))
                values[day] = ("NA" if measurement and measurement.root_unavailable else root,
                               "NA" if measurement and measurement.shoot_unavailable else shoot)
            wide_sheet.append((summary_number, source_name, original_number, slot.dish_number, slot.seedling_number,
                position, taxon.common_name, taxon.scientific_name) +
                tuple(value for day in all_dag for value in values.get(day, (None, None))) +
                (germinated_text, "是" if slot.sample_id is not None else "否", slot.planned_number, scope_text, slot.replicate_no))
    explanation.append(("项目", "说明"))
    notes = [
        ("导出时间", datetime.now(timezone.utc).isoformat()),
        ("软件版本", VERSION),
        ("所选实验", "；".join(f"{experiments[item].name}（{experiments[item].code}）" for item in experiment_ids)),
        ("材料数量", len(rows)),
        ("排序规则", "按中文名完整拼音、学名、原始材料编号、系统批次编号排序"),
        ("汇总编号", "仅属于本次导出，按排序后的实验材料从 001 编起，不写回实验数据"),
        ("原实验编号", "实验内确认置床编号时固定的编号；不同实验可各自从 001 开始"),
        ("培养皿现场编号", "单重复使用实验内材料编号，多重复在编号后加 -1、-2 等；不导出内部技术编号"),
        ("幼苗编号", "培养皿现场编号加皿内幼苗序号，例如单重复 001-01，多重复 001-1-01"),
        ("DAG", "发芽后测定时间，幼苗实际发芽后第 N 天"),
        ("测定值含义", "0 是实测零值；NA 表示已有记录中明确无法测量；空白表示没有测定值，不自动解释为漏测或最终未获得幼苗。根长、苗长单位均为 mm"),
        ("计划测定日期", "以幼苗发芽判定时间的实验室本地日期加 DAG 自然日计算；延迟天数按实际测定日期计算"),
        ("测定时间时区", "幼苗测定长表的发芽判定时间和实际测定时间按系统配置的实验室时区显示，并带时区偏移"),
        ("发芽率汇总", "仅以实际已置床的培养皿种子数为分母；没有巡检记录时累计发芽数和发芽率留空，明确记录 0 才表示已巡检且没有发芽"),
        ("测定数据", "宽表每行是一个计划幼苗槽位；长表每行是该槽位在本实验配置的一个发芽后测定时间。全部材料保留，空计划槽位只用于导出，不创建实际幼苗或测定记录"),
        ("数据状态", "已测定：已有测定记录；无测定记录：已有实际幼苗但该时间点没有记录；无实际幼苗：计划槽位尚无实际幼苗。仅表达当前事实，不判断实验最终结果"),
        ("计划取样序号", "按实验方案与材料覆盖值生成。每皿取样按各重复展开；每材料取样在全部重复间共用取样数，按实际发芽时间、重复和幼苗序号匹配，实际幼苗编号不改变"),
        ("计划培养皿重复", "按每皿取样时标明计划重复，尚未建立或已取消的培养皿仍保留设计位置；跨皿共用取样数时留空，不预先分配实际培养皿"),
        ("跨皿取样空槽位", "每材料取样且有多个重复时，尚无实际幼苗的槽位不指定培养皿，培养皿和幼苗编号留空；使用原实验编号与计划取样序号识别槽位"),
        ("联合导出时间点", "宽表列汇集所选实验的测定时间；长表仅展开各实验自身配置的时间点，其他实验独有的时间点在宽表留空"),
    ]
    for note in notes:
        explanation.append(note)
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
