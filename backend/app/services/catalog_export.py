"""User-facing catalog ledger; identifiers are written as text, never formulas."""
from io import BytesIO
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.models import Taxon, SeedLot
from app.services.ordering import taxon_key, material_key

TAXON_HEADERS = ("系统物种编号", "中文名", "学名", "科", "属", "生活型", "状态", "备注")
LOT_HEADERS = ("系统批次编号", "系统物种编号", "中文名", "学名", "原始材料编号", "来源", "采集/获得日期", "数量（粒）", "状态", "备注")


def append_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'
            cell.number_format = '@'


def build(db: Session) -> BytesIO:
    workbook = Workbook()
    taxa = workbook.active
    taxa.title = "01_物种"
    lots = workbook.create_sheet("02_种子批次")
    append_row(taxa, TAXON_HEADERS)
    append_row(lots, LOT_HEADERS)
    for item in sorted(db.scalars(select(Taxon)).all(), key=taxon_key):
        append_row(taxa, (item.code, item.common_name, item.scientific_name, item.family,
                         item.genus, item.life_form, "使用中" if item.is_active else "已停用", item.notes))
    rows = db.execute(select(SeedLot, Taxon).join(Taxon, SeedLot.taxon_id == Taxon.id)).all()
    for lot, taxon in sorted(rows, key=lambda pair: material_key(pair[1], pair[0])):
        append_row(lots, (lot.code, taxon.code, taxon.common_name, taxon.scientific_name,
                         lot.source_code, lot.source, lot.collected_at.isoformat() if lot.collected_at else None,
                         lot.quantity, "使用中" if lot.is_active else "已停用", lot.notes))
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
