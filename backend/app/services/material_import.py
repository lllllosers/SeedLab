"""Preview and atomically import a complete material inventory."""

from collections import defaultdict
from datetime import date, datetime, time
import unicodedata
from hashlib import sha256
from io import BytesIO

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Experiment, ExperimentMaterial, ImportJob, SeedLot, Taxon
from app.services.common import commit_or_conflict, flush_or_conflict, next_code, record


COLUMNS = ("系统物种编号", "中文名", "学名", "科", "属", "生活型", "原始材料编号",
           "来源", "采集/获得日期", "数量（粒）", "物种备注", "种子批次备注")
TAXON_FIELDS = {"中文名": "common_name", "科": "family", "属": "genus",
                "生活型": "life_form", "物种备注": "notes"}
LOT_FIELDS = {"来源": "source", "采集/获得日期": "collected_at",
              "数量（粒）": "quantity", "种子批次备注": "notes"}


def _text(value) -> str | None:
    result = str(value).strip() if value is not None else ""
    return result or None


def _date(value, line: int):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    try:
        return datetime.fromisoformat(str(value).strip()).replace(tzinfo=None)
    except ValueError as exc:
        raise ValueError(f"第 {line} 行：采集/获得日期请填写有效日期") from exc


def _quantity(value, line: int):
    if value is None or str(value).strip() == "":
        return None
    try:
        number = int(value)
        if isinstance(value, bool) or number < 0 or number > 2**63 - 1 or (
                isinstance(value, float) and value != number):
            raise ValueError
        return number
    except (ValueError, TypeError) as exc:
        raise ValueError(f"第 {line} 行：数量（粒）请填写非负整数或留空") from exc


def parse_file(content: bytes) -> list[dict]:
    if len(content) > 8_000_000:
        raise HTTPException(413, "文件不能超过 8 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            header = [_text(value) or "" for value in next(rows)]
            if header != list(COLUMNS):
                raise ValueError("表头与下载的种子材料模板不一致，请使用最新模板")
            parsed = []
            for line, values in enumerate(rows, start=2):
                if all(value is None or str(value).strip() == "" for value in values):
                    continue
                if len(parsed) >= 500:
                    raise ValueError("一次最多预检 500 行材料")
                raw = dict(zip(COLUMNS, values))
                row = {column: _text(raw.get(column)) for column in COLUMNS}
                row_errors = []
                for column, parser in (("采集/获得日期", _date), ("数量（粒）", _quantity)):
                    try:
                        row[column] = parser(raw.get(column), line)
                    except ValueError as exc:
                        row[column] = None
                        row_errors.append(str(exc))
                row["_errors"] = row_errors
                row["line"] = line
                parsed.append(row)
            if not parsed:
                raise ValueError("文件没有材料数据行，请在模板中填写至少一份材料")
            return parsed
        finally:
            workbook.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"无法读取种子材料文件：{exc}") from exc


def chinese_comparison_key(value: str | None) -> str:
    """Comparison only: keep punctuation, parenthesized numbers and display values."""
    return " ".join(unicodedata.normalize("NFKC", value or "").split())


def scientific_comparison_key(value: str | None) -> str:
    value = unicodedata.normalize("NFKC", value or "").replace("_", " ")
    return " ".join(value.split()).casefold()


def _lot_changes(row: dict, lot: SeedLot) -> tuple[list[str], list[str]]:
    fills, conflicts = [], []
    for label, field in LOT_FIELDS.items():
        value, current = row[label], getattr(lot, field)
        if value is not None:
            if current is None or current == "":
                fills.append(label)
            elif current != value:
                conflicts.append(label)
    return fills, conflicts


def _ger_history(db):
    """One ID-based join; a species includes every lot, even a different seed source."""
    taxa, lots = defaultdict(set), defaultdict(set)
    query = (select(SeedLot.taxon_id, SeedLot.id, Experiment.status)
             .join(ExperimentMaterial, ExperimentMaterial.seed_lot_id == SeedLot.id)
             .join(Experiment, Experiment.id == ExperimentMaterial.experiment_id)
             .where(Experiment.experiment_type == "GER"))
    for taxon_id, lot_id, state in db.execute(query):
        taxa[taxon_id].add(state)
        lots[lot_id].add(state)
    return taxa, lots


def _ger_status(states):
    return ("completed" if "completed" in states else "active" if "active" in states
            else "recorded" if states else "none")


def _taxon_match(row, taxa, choice):
    code, sci, common = row["系统物种编号"], scientific_comparison_key(row["学名"]), chinese_comparison_key(row["中文名"])
    errors, warnings = [], []
    if code:
        found = next((t for t in taxa if t.code == code), None)
        if not found:
            return "error", None, [], [f"第 {row['line']} 行填写的物种编号 {code} 不存在，请核对"], warnings
        if sci and sci != scientific_comparison_key(found.scientific_name):
            errors.append(f"第 {row['line']} 行：物种编号与学名不对应，请核对")
        candidates = [found]
    else:
        candidates = [t for t in taxa if sci and scientific_comparison_key(t.scientific_name) == sci]
    exact = bool(candidates)
    if not candidates:
        candidates = [t for t in taxa if common and chinese_comparison_key(t.common_name) == common]
    found = candidates[0] if exact and len(candidates) == 1 else None
    status = "existing" if found else "confirm" if candidates else "new"
    if status == "confirm" and choice is not None:
        if choice == "new":
            if exact:
                errors.append("系统中已有学名格式一致的物种，请明确选择要复用的物种，不能重复新建。")
            else:
                status = "new"
        elif isinstance(choice, str) and choice.startswith("reuse:"):
            found = next((t for t in candidates if t.id == choice.removeprefix("reuse:")), None)
            if found:
                status = "existing"
            else:
                errors.append("所选物种不在候选范围内，请重新预检并选择。")
        else:
            errors.append("请明确选择复用已有物种或确认为新物种。")
    elif choice is not None and ((choice != "new" and status == "new") or
            (status == "existing" and choice != f"reuse:{found.id}")):
        errors.append("物种识别结果已变化，请重新预检后选择。")
    if status == "new" and not sci:
        errors.append(f"第 {row['line']} 行：新物种必须填写学名；已有物种请选择复用。")
    if found:
        if common and common != chinese_comparison_key(found.common_name):
            warnings.append("中文名与系统记录不同，请选择沿用系统名称或更新名称。")
        if row["学名"] and sci != scientific_comparison_key(found.scientific_name):
            warnings.append(f"按您的选择复用系统物种；保留已有学名 {found.scientific_name}。")
    elif status == "confirm":
        warnings.append("系统中有同名或学名格式一致的物种，请核对系统学名与表中学名后明确选择；不会自动判断学名同义关系。")
    return "error" if errors else status, found, candidates, errors, warnings


def _material_match(row, lots, mode):
    """Resolve lot identity only after species identity; never merge on species alone."""
    code = row["原始材料编号"]
    exact = [lot for lot in lots if code and lot.source_code == code]
    if code:
        return ("existing", exact[0], exact) if len(exact) == 1 else ("confirm", None, exact) if exact else ("new", None, [])
    if not lots:
        return "new", None, []
    exact = [lot for lot in lots if row["来源"] and row["采集/获得日期"]
             and lot.source == row["来源"] and lot.collected_at == row["采集/获得日期"]]
    if len(exact) == 1:
        return ("existing" if mode == "complete" else "confirm"), exact[0], exact
    if len(exact) > 1:
        return "confirm", None, exact
    fields = (("来源", "source"), ("采集/获得日期", "collected_at"))
    compatible = [lot for lot in lots if not any(row[label] is not None and getattr(lot, field) is not None
                  and row[label] != getattr(lot, field) for label, field in fields)]
    if not compatible:
        return "new", None, []
    active = [lot for lot in lots if lot.is_active]
    if (mode == "complete" and any(row[label] is not None for label, _ in fields)
            and len(active) == 1 and active[0] in compatible):
        return "existing", active[0], compatible
    return "confirm", None, compatible


def experiment_selection_summary(rows, decisions=None):
    """Project experimental populations; import identity decisions remain explicit."""
    decisions = decisions or {}
    species, batches, pending_taxa, pending_materials = {}, {}, set(), set()
    for row in rows:
        row['selection_species_group'] = None
        row['selection_batch_group'] = None
        if row['errors']:
            continue
        if row['taxon_status'] == 'confirm':
            pending_taxa.add(row['line'])
            continue
        species_key = row['taxon_id'] or 'new:' + scientific_comparison_key(row['scientific_name'])
        if row['taxon_ger_status'] == 'none':
            species.setdefault(species_key, []).append(row['line'])
            row['selection_species_group'] = species_key
        choice = decisions.get(str(row['line'])) if row['material_status'] == 'confirm' else None
        status = row['material_status']
        if status == 'confirm' and not _valid_decision(row, choice):
            pending_materials.add(row['line'])
        if choice and choice.startswith('reuse:'):
            chosen = next((lot for lot in row['candidate_lots'] if choice == 'reuse:' + lot['id']), None)
            needs = bool(chosen and chosen['ger_status'] == 'none')
            batch_key = 'lot:' + chosen['id'] if chosen else None
        elif status in {'existing', 'updatable'}:
            needs = row['material_ger_status'] == 'none'
            batch_key = 'lot:' + str(row['candidate_lot_id'])
        else:
            same_history = bool(not row['source_code'] and row['source'] and row['source'] in row['historical_sources'])
            needs = choice == 'new' or not same_history
            batch_key = json_batch_key(species_key, row)
            if choice == 'new' and row['candidate_row']:
                batch_key += ':separate:' + str(row['line'])
            if choice and choice.startswith('reuse-row:'):
                previous = next((r for r in rows if str(r['line']) == choice.removeprefix('reuse-row:')), None)
                batch_key = previous.get('selection_batch_group') if previous else None
                needs = batch_key is not None
        if needs and batch_key:
            batches.setdefault(batch_key, []).append(row['line'])
            row['selection_batch_group'] = batch_key
    return {'new_species_unique_count': len(species), 'new_batch_unique_count': len(batches),
            'pending_taxon_count': len(pending_taxa), 'pending_material_count': len(pending_materials),
            'species_groups': list(species.values()), 'batch_groups': list(batches.values())}


def json_batch_key(species_key, row):
    from json import dumps
    identity = [species_key, row['source_code']] if row['source_code'] else [species_key, row['source'], row.get('collected_at')]
    return dumps(identity, ensure_ascii=False)


def preview(db: Session, content: bytes, mode: str = "complete", taxon_decisions: dict | None = None,
            decisions: dict | None = None) -> dict:
    if mode not in {"complete", "new"}:
        raise HTTPException(422, "请选择完整清单或本次新增材料，再重新预检")
    taxon_decisions = taxon_decisions or {}
    file_hash = sha256(content).hexdigest()
    already = db.scalar(select(ImportJob.id).where(ImportJob.file_hash == file_hash, ImportJob.status == "completed")) is not None
    taxa = list(db.scalars(select(Taxon).order_by(Taxon.code)))
    lots = list(db.scalars(select(SeedLot).order_by(SeedLot.code)))
    taxa_history, lot_history = _ger_history(db)
    results, seen_identity = [], {}
    for row in parse_file(content):
        line = row["line"]
        taxon_status, taxon, candidates, errors, warnings = _taxon_match(row, taxa, taxon_decisions.get(str(line)))
        errors = list(row["_errors"]) + errors
        fills, conflicts = [], []
        if taxon:
            for label, field in TAXON_FIELDS.items():
                value, current = row[label], getattr(taxon, field)
                if value is not None:
                    if current is None or current == "":
                        fills.append(label)
                    elif current != value:
                        conflicts.append(label)
        material_status, candidate, possible = ("pending", None, []) if taxon_status in {"confirm", "error"} else _material_match(
            row, [lot for lot in lots if taxon and lot.taxon_id == taxon.id], mode)
        candidate_row = None
        if material_status != "pending":
            taxon_key = taxon.id if taxon else "new:" + scientific_comparison_key(row["学名"])
            identity = ((taxon_key, row["原始材料编号"]) if row["原始材料编号"] else
                        (taxon_key, row["来源"], row["采集/获得日期"]))
            if identity in seen_identity:
                candidate_row = seen_identity[identity]
                material_status = "confirm"
            else:
                seen_identity[identity] = line
        if candidate:
            lot_fills, lot_conflicts = _lot_changes(row, candidate)
            fills.extend(lot_fills)
            conflicts.extend(lot_conflicts)
            if material_status == "existing" and (fills or conflicts):
                material_status = "updatable"
        overall = "error" if errors else "confirm" if taxon_status == "confirm" or material_status == "confirm" else material_status
        if errors:
            material_status = "error"
        taxon_states = taxa_history.get(taxon.id, set()) if taxon else set()
        material_states = lot_history.get(candidate.id, set()) if candidate else set()
        results.append({"line": line, "status": "registered" if overall == "existing" else overall,
            "taxon_status": taxon_status, "material_status": material_status,
            "overall_action_required": overall in {"error", "confirm", "pending"},
            "taxon_id": taxon.id if taxon else None, "new_taxon": taxon_status == "new",
            "taxon_candidates": [{"id": t.id, "code": t.code, "common_name": t.common_name,
                "scientific_name": t.scientific_name, "ger_status": _ger_status(taxa_history.get(t.id, set()))} for t in candidates],
            "taxon_ger_status": "pending" if taxon_status == "confirm" else _ger_status(taxon_states),
            "taxon_ger_states": sorted(taxon_states),
            "material_ger_status": "pending" if material_status in {"confirm", "pending"} else _ger_status(material_states),
            "material_ger_states": sorted(material_states),
            "candidate_lot_id": candidate.id if candidate else None,
            "candidate_lots": [{"id": lot.id, "code": lot.code, "is_active": lot.is_active, "source_code": lot.source_code,
                "source": lot.source, "collected_at": lot.collected_at.isoformat() if lot.collected_at else None,
                "ger_status": _ger_status(lot_history.get(lot.id, set()))} for lot in possible],
            "candidate_row": candidate_row, "name": row["中文名"] or row["学名"], "source": row["来源"],
            "collected_at": row["采集/获得日期"].isoformat() if row["采集/获得日期"] else None,
            "system_common_name": taxon.common_name if taxon else None,
            "name_difference": bool(taxon and row["中文名"] and chinese_comparison_key(row["中文名"]) != chinese_comparison_key(taxon.common_name)),
            "can_update_common_name": bool(taxon and scientific_comparison_key(row["学名"]) == scientific_comparison_key(taxon.scientific_name)),
            "historical_sources": [lot.source for lot in lots if taxon and lot.taxon_id == taxon.id and lot_history.get(lot.id)],
            "scientific_name": row["学名"] or (taxon.scientific_name if taxon else None),
            "source_code": row["原始材料编号"], "fills": fills, "conflicts": conflicts, "warnings": warnings,
            "suggested_decision": f"reuse-row:{candidate_row}" if candidate_row else f"reuse:{candidate.id}" if candidate else None,
            "errors": errors})
    stats = {"total": len(results), **{state: sum(r["status"] == state for r in results)
             for state in ("registered", "new", "updatable", "confirm", "error")}}
    stats["errors"] = stats.pop("error")
    stats["new_taxa"] = len({scientific_comparison_key(r["scientific_name"]) for r in results if r["taxon_status"] == "new" and not r["errors"]})
    stats["reused_taxa"] = sum(r["taxon_status"] == "existing" for r in results)
    for prefix, field, states in (("taxon", "taxon_status", ("existing", "new", "confirm")),
                                  ("material", "material_status", ("existing", "new", "confirm", "updatable", "pending"))):
        stats.update({f"{prefix}_{state}": sum(r[field] == state for r in results) for state in states})
    stats["taxon_with_ger"] = sum(r["taxon_ger_status"] in {"completed", "active", "recorded"} for r in results)
    stats["taxon_without_ger"] = sum(r["taxon_ger_status"] == "none" for r in results)
    selection = experiment_selection_summary(results, decisions)
    return {"file_hash": file_hash, "already_imported": already, "mode": mode, "rows": results, "stats": stats,
            "experiment_selection_summary": selection}


def _valid_decision(row: dict, choice: str | None) -> bool:
    if not isinstance(choice, str):
        return False
    if choice == "new":
        return True
    if choice == "reuse":
        return row["candidate_row"] is not None or row["candidate_lot_id"] is not None
    if choice.startswith("reuse-row:"):
        return row["candidate_row"] is not None and choice == f"reuse-row:{row['candidate_row']}"
    if choice.startswith("reuse:"):
        return choice.removeprefix("reuse:") in {lot["id"] for lot in row["candidate_lots"]}
    return False


def confirm(db: Session, content: bytes, filename: str, decisions: dict[str, str], user_id: str,
            mode: str = "complete", taxon_decisions: dict | None = None, name_decisions: dict | None = None) -> dict:
    report = preview(db, content, mode, taxon_decisions)
    if report["already_imported"]:
        raise HTTPException(409, "这份文件已经成功导入过；请使用更新后的完整清单")
    if report["stats"]["errors"]:
        raise HTTPException(422, "文件仍有错误，请先按预检提示修正后重新上传")
    if any(r["taxon_status"] == "confirm" for r in report["rows"]):
        raise HTTPException(422, "请先为所有需要确认的物种选择复用或新增，并重新预检材料。")
    if any(r["material_status"] == "confirm" and not _valid_decision(r, decisions.get(str(r["line"]))) for r in report["rows"]):
        raise HTTPException(422, "请先为所有需要确认的材料选择复用或新增")
    by_name = defaultdict(list)
    for taxon in db.scalars(select(Taxon)):
        by_name[scientific_comparison_key(taxon.scientific_name)].append(taxon)
    created_by_line, all_ids, created_ids, updated_ids = {}, set(), set(), set()
    new_taxa = new_lots = 0
    taxon_next = int(next_code(db, Taxon, "SP-").removeprefix("SP-"))
    lot_prefix = f"LOT-{datetime.now().year}-"
    lot_next = int(next_code(db, SeedLot, lot_prefix, 3).rsplit("-", 1)[1])
    try:
        renames = {}
        for result in report['rows']:
            choice = (name_decisions or {}).get(str(result['line']), 'keep')
            if not isinstance(choice, str) or choice not in {'keep', 'update'}:
                raise HTTPException(422, '名称处理选项无效，请重新预检。')
            if choice == 'update':
                if not result['name_difference'] or not result['can_update_common_name']:
                    raise HTTPException(422, '仅可为已确认学名相同的已有物种更新中文名，请重新预检。')
                taxon_id, name = result['taxon_id'], result['name']
                if taxon_id in renames and renames[taxon_id] != name:
                    raise HTTPException(422, '同一物种选择了不同的新中文名，请统一选择后再导入。')
                renames[taxon_id] = name
        for taxon_id, name in renames.items():
            taxon = db.get(Taxon, taxon_id)
            before = {'common_name': taxon.common_name}
            taxon.common_name = name
            record(db, user_id, 'update', 'Taxon', taxon.id, before, {'common_name': name})
        for row, result in zip(parse_file(content), report["rows"]):
            line = row["line"]
            taxon = db.get(Taxon, result["taxon_id"]) if result["taxon_id"] else None
            key = scientific_comparison_key(row["学名"])
            if taxon is None and by_name[key]:
                # Earlier rows in this same transaction may have created the species.
                if len(by_name[key]) != 1:
                    raise HTTPException(422, "同一学名对应多个物种，请重新预检并明确选择。")
                taxon = by_name[key][0]
            taxon_updated = taxon is not None and taxon.id in renames
            if taxon is None:
                taxon = Taxon(code=f"SP-{taxon_next:04d}", scientific_name=row["学名"],
                              **{field: row[label] for label, field in TAXON_FIELDS.items()})
                taxon_next += 1
                db.add(taxon)
                flush_or_conflict(db)
                by_name[key].append(taxon)
                new_taxa += 1
                record(db, user_id, "import", "Taxon", taxon.id, None, {"code": taxon.code, "line": line})
            else:
                before, after = {}, {}
                for label, field in TAXON_FIELDS.items():
                    value = row[label]
                    if value is not None and not getattr(taxon, field):
                        before[field] = None
                        setattr(taxon, field, value)
                        after[field] = value
                if after:
                    taxon_updated = True
                    record(db, user_id, "update", "Taxon", taxon.id, before, after)
            decision = decisions.get(str(line))
            lot = db.get(SeedLot, result["candidate_lot_id"]) if result["candidate_lot_id"] else None
            if result["material_status"] == "confirm":
                if decision.startswith("reuse:"):
                    lot = db.get(SeedLot, decision.removeprefix("reuse:"))
                elif decision.startswith("reuse-row:"):
                    lot = created_by_line.get(int(decision.removeprefix("reuse-row:")))
                elif decision == "reuse" and result["candidate_row"] is not None:
                    lot = created_by_line.get(result["candidate_row"])
            if result["material_status"] == "new" or (result["material_status"] == "confirm" and decision == "new"):
                lot = SeedLot(code=f"{lot_prefix}{lot_next:03d}", taxon_id=taxon.id, source_code=row["原始材料编号"],
                              **{field: row[label] for label, field in LOT_FIELDS.items()})
                lot_next += 1
                db.add(lot)
                flush_or_conflict(db)
                new_lots += 1
                created_ids.add(lot.id)
                record(db, user_id, "import", "SeedLot", lot.id, None, {"code": lot.code, "line": line})
            elif lot is not None and lot.taxon_id == taxon.id:
                before, after = {}, {}
                for label, field in LOT_FIELDS.items():
                    value = row[label]
                    if value is not None and getattr(lot, field) in (None, ""):
                        before[field] = getattr(lot, field)
                        setattr(lot, field, value)
                        after[field] = value.isoformat() if isinstance(value, datetime) else value
                if after or taxon_updated:
                    updated_ids.add(lot.id)
                if after:
                    record(db, user_id, "update", "SeedLot", lot.id, before, after)
            else:
                raise HTTPException(422, f"第 {line} 行找不到该物种下可复用的材料，请重新预检")
            all_ids.add(lot.id)
            created_by_line[line] = lot
        job = ImportJob(user_id=user_id, filename=filename[:255], status="completed",
                        total_rows=len(report["rows"]), successful_rows=len(report["rows"]), file_hash=report["file_hash"])
        db.add(job)
        flush_or_conflict(db)
        record(db, user_id, "import", "ImportJob", job.id, None,
               {"filename": job.filename, "total_rows": job.total_rows, "new_taxa": new_taxa, "new_lots": new_lots})
        commit_or_conflict(db)
    except Exception:
        db.rollback()
        raise
    return {"job_id": job.id, "total": job.total_rows, "new_taxa": new_taxa, "new_lots": new_lots,
            "all_seed_lot_ids": sorted(all_ids), "created_seed_lot_ids": sorted(created_ids),
            "total_material_count": len(all_ids), "existing_material_count": len(all_ids - created_ids),
            "created_material_count": len(created_ids), "updated_material_count": len(updated_ids - created_ids)}
