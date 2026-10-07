"""Public GER flow and retired importer boundary, using isolated test databases."""
import ast
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import pytest

from app.analysis import read_dataset
from app.db.session import get_db
from app.services.measurement_slots import MeasurementDatasetReader
from test_architecture_dependencies import assert_dependencies
from test_lifecycle_contract import audit_rows
from test_seedling_measurement import payload


ROOT = Path(__file__).resolve().parents[2]
LEGACY_MODULES = ('scripts.maintenance', 'backfill_legacy_200_species', 'legacy_importer')
LEGACY_TOKENS = ('backfill_legacy_200_species', 'seedlablegacyimport', 'legacy_importer')


def assert_no_legacy_runtime(source, module):
    assert_dependencies(source, module, LEGACY_MODULES)
    # Also reject literal dynamic imports and process launches of the old importer.
    for node in ast.walk(ast.parse(source)):
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            assert not any(token in node.value.casefold() for token in LEGACY_TOKENS), \
                f'{module}: legacy importer runtime reference'


def test_product_runtime_and_packaging_do_not_depend_on_legacy_importer():
    files = list((ROOT / 'backend/app').rglob('*.py')) + list((ROOT / 'control_center').rglob('*.py'))
    files += [ROOT / name for name in ('packaging/server_entry.py', 'packaging/control_center_entry.py',
              'packaging/seedlab.spec', 'scripts/run_prod.py', 'scripts/run_control_center.py')]
    for path in files:
        module = path.with_suffix('').relative_to(ROOT).as_posix().replace('/', '.')
        if module.startswith('backend.'):
            module = module.removeprefix('backend.')
        assert_no_legacy_runtime(path.read_text(encoding='utf-8'), module)


@pytest.mark.parametrize('source', [
    'from scripts.maintenance.backfill_legacy_200_species import apply',
    'import backfill_legacy_200_species',
    'from legacy_importer import run',
    'importlib.import_module("scripts.maintenance.backfill_legacy_200_species")',
    '__import__("legacy_importer")',
    'subprocess.run(["SeedLabLegacyImport-v1.0.0.exe"])',
])
def test_runtime_guard_rejects_injected_legacy_references(source):
    with pytest.raises(AssertionError):
        assert_no_legacy_runtime(source, 'app.main')


def test_normal_material_import_and_maintenance_assets_remain_distinct():
    assert_no_legacy_runtime('from app.services.material_import import preview', 'app.api.material_import')
    assert (ROOT / 'scripts/maintenance/backfill_legacy_200_species.py').is_file()
    assert (ROOT / 'backend/tests/test_legacy_backfill.py').is_file()


def create_design(client, headers, lots, name):
    created = client.post('/api/experiments', headers=headers,
                          json={'experiment_type': 'GER', 'name': name})
    assert created.status_code == 201, created.text
    experiment = created.json()
    base = f"/api/experiments/{experiment['id']}"
    protocol = {'seeds_per_dish': 10, 'replicate_count': 1, 'observation_period_days': None,
                'sampling_rule': 'first_germinated', 'sample_count': 2, 'sample_scope': 'per_dish',
                'germination_criterion': '胚根露出'}
    assert client.put(f'{base}/protocol', headers=headers, json=protocol).status_code == 200
    available = client.get('/api/experiments/available-seed-lots').json()
    for lot in lots:
        assert any(row['id'] == lot['id'] for row in available)
        added = client.post(f'{base}/materials', headers=headers, json={'seed_lot_id': lot['id']})
        assert added.status_code == 201, added.text
    assert client.put(f'{base}/dag', headers=headers, json={'days': [0, 1]}).status_code == 200
    assert client.patch(base, headers=headers, json={'status': 'ready'}).status_code == 200
    confirmed = client.post(f'{base}/confirm-numbers', headers=headers)
    assert confirmed.status_code == 200, confirmed.text
    dishes = confirmed.json()['dishes']
    assert all(dish['sown_at'] is None for dish in dishes)
    return base, experiment, dishes, protocol


def new_lot(client, headers, scientific, common):
    taxon = client.post('/api/taxa', headers=headers,
                        json={'scientific_name': scientific, 'common_name': common})
    assert taxon.status_code == 201, taxon.text
    lot = client.post('/api/seed-lots', headers=headers, json={'taxon_id': taxon.json()['id']})
    assert lot.status_code == 201, lot.text
    return lot.json()


def test_complete_public_ger_flow_keeps_plans_facts_zero_na_and_correction(auth_client):
    client, headers = auth_client  # Fixture performs real login and CSRF exchange.
    assert client.get('/api/experiments/types').json() == [{'value': 'GER', 'label': '种子萌发试验'}]
    lots = [new_lot(client, headers, 'Setaria viridis', '狗尾草'),
            new_lot(client, headers, 'Poa annua', '早熟禾')]
    base, experiment, dishes, _ = create_design(client, headers, lots, '完整萌发流程验证')
    now = datetime.now(timezone.utc)
    sown = client.post(f'{base}/sowing/batch', headers=headers,
        json={'dish_ids': [dish['id'] for dish in dishes], 'sown_at': (now - timedelta(days=5)).isoformat()})
    assert sown.status_code == 200 and sown.json()['experiment']['status'] == 'active'
    assert client.post(f'{base}/reopen-design', headers=headers).status_code == 409
    when = now - timedelta(days=3)
    observation = client.post(f'{base}/observations/batch', headers=headers,
        json={'observed_at': when.isoformat(), 'entries': [
            {'dish_id': dishes[0]['id'], 'new_germinated_count': 5},
            {'dish_id': dishes[1]['id'], 'new_germinated_count': 0}]})
    assert observation.status_code == 200, observation.text
    samples = client.get(f'{base}/samples').json()
    assert len(samples) == 2 and {row['sample_number'] for row in samples} == {1, 2}
    assert all(row['germinated_at'] == when.isoformat().replace('+00:00', 'Z') for row in samples)
    tasks = client.get(f'{base}/measurement-tasks').json()['tasks']
    assert len(tasks) == 4 and {row['day_after_germination'] for row in tasks} == {0, 1}
    with contextmanager(client.app.dependency_overrides[get_db])() as db:
        before = read_dataset(MeasurementDatasetReader(db), [experiment['id']])
    assert len(before.seedling_slots) == 4 and len(before.rows) == 8
    assert sum(row.slot.sample_id is None for row in before.rows) == 4
    assert sum(row.data_status == '无测定记录' for row in before.rows) == 4
    saved = []
    for index, task in enumerate(tasks):
        values = {'root': 0, 'shoot': None, 'shoot_na': True} if index == 0 else \
                 {'root': None, 'root_na': True, 'shoot': 1.25} if index == 1 else {'root': 2, 'shoot': 3}
        response = client.post(f'{base}/measurements', headers=headers, json=payload(task, now, **values))
        assert response.status_code == 201, response.text
        saved.append(response.json())
    assert client.get(f'{base}/completion-check').json()['can_complete'] is True
    assert client.post(f'{base}/complete', headers=headers).json()['status'] == 'completed'
    correction = client.patch(f"{base}/measurements/{saved[0]['id']}", headers=headers,
                              json={'notes': '完成后复核'})
    assert correction.status_code == 200 and correction.json()['root_length_mm'] == 0
    assert correction.json()['shoot_unavailable'] is True
    assert client.delete(f"{base}/measurements/{saved[0]['id']}", headers=headers).status_code == 409
    with contextmanager(client.app.dependency_overrides[get_db])() as db:
        dataset = read_dataset(MeasurementDatasetReader(db), [experiment['id']])
    assert sum(row.measurement_exists for row in dataset.rows) == 4
    assert any(row.root_length_mm == Decimal(0) and row.shoot_unavailable for row in dataset.rows)
    assert any(row.root_unavailable and row.shoot_length_mm == Decimal('1.25') for row in dataset.rows)
    assert all(row.root_unavailable is None for row in dataset.rows if row.slot.sample_id is None)
    output = client.post('/api/export/experiments/workbook.xlsx', headers=headers,
                         json={'experiment_ids': [experiment['id']]})
    assert output.status_code == 200
    book = load_workbook(BytesIO(output.content), data_only=True)
    try:
        assert book.sheetnames == ['01_材料总表', '02_发芽率汇总', '03_发芽原始记录',
                                   '04_幼苗测定长表', '05_幼苗测定宽表', '06_导出说明']
        long = list(book['04_幼苗测定长表'].values)
        wide = list(book['05_幼苗测定宽表'].values)
        assert len(long) == 9 and len(wide) == 5
        rows = [dict(zip(long[0], row)) for row in long[1:]]
        assert any(row['根长（mm）'] == 0 and row['苗长状态'] == '无法测量' for row in rows)
        assert any(row['根长状态'] == '无法测量' and row['苗长（mm）'] == 1.25 for row in rows)
        assert sum(row['根长（mm）'] is None and row['苗长（mm）'] is None for row in rows) == 4
    finally:
        book.close()
    assert any(row['after'].get('status') == 'completed' for row in audit_rows(client, 'Experiment', experiment['id']))
    assert any(row['action'] == 'update' for row in audit_rows(client, 'SeedlingMeasurement', saved[0]['id']))


def test_reopen_numbering_then_terminate_preserves_facts_and_export(auth_client):
    client, headers = auth_client
    lot = new_lot(client, headers, 'Poa pratensis', '草地早熟禾')
    base, experiment, _, protocol = create_design(client, headers, [lot], '编号与终止验证')
    assert client.patch(base, headers=headers, json={'status': 'draft'}).status_code == 409
    reopened = client.post(f'{base}/reopen-design', headers=headers)
    assert reopened.status_code == 200
    assert reopened.json()['experiment']['numbering_locked_at'] is None
    assert reopened.json()['dishes'] == []
    assert client.get(base).json()['code'] == experiment['code']
    assert client.put(f'{base}/protocol', headers=headers, json=protocol).status_code == 200
    assert client.put(f'{base}/dag', headers=headers, json={'days': [0, 1]}).status_code == 200
    assert client.patch(base, headers=headers, json={'status': 'ready'}).status_code == 200
    dish = client.post(f'{base}/confirm-numbers', headers=headers).json()['dishes'][0]
    now = datetime.now(timezone.utc)
    assert client.post(f'{base}/sowing/batch', headers=headers,
        json={'dish_ids': [dish['id']], 'sown_at': (now - timedelta(days=2)).isoformat()}).status_code == 200
    assert client.post(f'{base}/observations/batch', headers=headers,
        json={'observed_at': (now - timedelta(days=1)).isoformat(),
              'entries': [{'dish_id': dish['id'], 'new_germinated_count': 1}]}).status_code == 200
    task = client.get(f'{base}/measurement-tasks').json()['tasks'][0]
    assert client.post(f'{base}/measurements', headers=headers, json=payload(task, now)).status_code == 201
    assert client.post(f'{base}/terminate', headers=headers, json={'reason': '材料污染'}).json()['status'] == 'cancelled'
    assert len(client.get(f'{base}/samples').json()) == 1
    assert client.get(f'{base}/measurement-worklist').json()['materials'] == []
    assert client.post('/api/export/experiments/workbook.xlsx', headers=headers,
        json={'experiment_ids': [experiment['id']]}).status_code == 200
    assert any(row['after'].get('termination_reason') == '材料污染'
               for row in audit_rows(client, 'Experiment', experiment['id']))
