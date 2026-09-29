from io import BytesIO
from datetime import datetime

from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app.db.session import make_engine


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["物种编号", "中文名", "学名", "来源", "数量", "备注"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload(client, headers, rows):
    return client.post("/api/import/seed-lots", headers=headers,
                       files={"file": ("seed-lots.xlsx", workbook_bytes(rows),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_seed_lot_import_200_rows_template_and_audit(auth_client, tmp_path):
    client, headers = auth_client
    taxa_template = client.get("/api/export/taxa-template.xlsx")
    assert taxa_template.status_code == 200
    assert list(load_workbook(BytesIO(taxa_template.content), read_only=True).active.values) == [
        ("学名", "中文名", "科")
    ]
    for index in range(200):
        response = client.post("/api/taxa", headers=headers,
                               json={"scientific_name": f"Species {index}", "common_name": f"物种{index}"})
        assert response.status_code == 201, response.text
    template = client.get("/api/export/seed-lots-template.xlsx")
    assert template.status_code == 200
    sheet = load_workbook(BytesIO(template.content), read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 201
    assert rows[0] == ("物种编号", "中文名", "学名", "来源", "数量", "备注")
    assert rows[1][:3] == ("SP-0001", "物种0", "Species 0")

    data_rows = [(*row[:3], "采集地", index, "人工测试") for index, row in enumerate(rows[1:])]
    imported = upload(client, headers, data_rows)
    assert imported.status_code == 201, imported.text
    assert imported.json()["imported"] == 200
    lots = client.get("/api/seed-lots").json()
    assert len(lots) == 200
    assert lots[0]["taxon"]["common_name"]
    assert lots[0]["taxon"]["scientific_name"]
    assert {item["code"] for item in lots} == {
        f"LOT-{datetime.now().year}-{index:03d}" for index in range(1, 201)
    }
    engine = make_engine(f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    with engine.connect() as connection:
        assert connection.execute(text("SELECT total_rows FROM import_jobs WHERE id=:id"),
                                  {"id": imported.json()["id"]}).scalar() == 200
        assert connection.execute(text("SELECT COUNT(*) FROM audit_logs WHERE action='import' AND entity_type='SeedLot'")).scalar() == 200
    engine.dispose()


def test_seed_lot_import_rejects_invalid_rows_without_partial_writes(auth_client, tmp_path):
    client, headers = auth_client
    first = client.post("/api/taxa", headers=headers,
                        json={"scientific_name": "Setaria viridis", "common_name": "狗尾草"}).json()
    inactive = client.post("/api/taxa", headers=headers,
                           json={"scientific_name": "Poa annua"}).json()
    client.patch(f"/api/taxa/{inactive['id']}", headers=headers, json={"is_active": False})
    rows = [(first["code"], "狗尾草", "Setaria viridis", "来源一", 5, None),
            ("SP-9999", None, None, "来源二", 2, None),
            (inactive["code"], None, None, None, 1, None),
            (first["code"], None, None, None, -1, None)]
    response = upload(client, headers, rows)
    assert response.status_code == 422
    assert all(f"第 {line} 行" in response.json()["detail"] for line in (3, 4, 5))
    assert "整表未导入" in response.json()["detail"]
    assert "第 3 行：找不到物种编号 SP-9999" in response.json()["detail"]
    assert "taxon_code" not in response.json()["detail"]
    assert client.get("/api/seed-lots").json() == []
    engine = make_engine(f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    with engine.connect() as connection:
        assert connection.execute(text("SELECT COUNT(*) FROM import_jobs")).scalar() == 0
        assert connection.execute(text("SELECT COUNT(*) FROM audit_logs WHERE entity_type='SeedLot'")).scalar() == 0
    engine.dispose()


def test_material_and_seed_lot_search_use_all_taxon_names(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers,
                        json={"scientific_name": "Setaria viridis", "common_name": "狗尾草"}).json()
    lot = client.post("/api/seed-lots", headers=headers,
                      json={"taxon_id": taxon["id"], "source": "野外采集"}).json()
    for term in ("狗尾草", "Setaria", taxon["code"]):
        available = client.get("/api/experiments/available-seed-lots", params={"q": term}).json()
        assert len(available) == 1
        assert available[0]["id"] == lot["id"]
        assert available[0]["taxon_common_name"] == "狗尾草"
        assert available[0]["taxon_scientific_name"] == "Setaria viridis"
        assert available[0]["taxon_code"] == taxon["code"]
        assert "taxon_name" not in available[0]
        listed = client.get("/api/seed-lots", params={"q": term}).json()
        assert [item["id"] for item in listed] == [lot["id"]]
    for term in (lot["code"], "野外采集"):
        assert [item["id"] for item in client.get("/api/seed-lots", params={"q": term}).json()] == [lot["id"]]


def test_taxa_chinese_template_rejects_invalid_sheet_without_partial_writes(auth_client):
    client, headers = auth_client
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["学名", "中文名", "科"])
    sheet.append(["Setaria viridis", "狗尾草", "禾本科"])
    sheet.append([None, "缺少学名", None])
    output = BytesIO()
    workbook.save(output)
    response = client.post("/api/import/taxa", headers=headers,
                           files={"file": ("taxa.xlsx", output.getvalue(),
                                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 422
    assert "第 3 行缺少学名" in response.json()["detail"]
    assert client.get("/api/taxa").json() == []


def test_seed_lot_import_nullable_quantity_and_new_lot_per_row(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers,
                        json={"scientific_name": "Setaria viridis", "common_name": "狗尾草"}).json()
    row = (taxon["code"], "狗尾草", "Setaria viridis", "", None, "")
    assert upload(client, headers, [row, row]).json()["imported"] == 2
    lots = client.get("/api/seed-lots").json()
    assert len(lots) == 2
    assert lots[0]["code"] != lots[1]["code"]
    assert all(item["quantity"] is None and item["taxon"]["common_name"] == "狗尾草" for item in lots)


def test_experiment_status_actions_keep_start_as_only_active_path(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers, json={"scientific_name": "Setaria viridis"}).json()
    lot = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxon["id"]}).json()
    experiment = client.post("/api/experiments/configured", headers=headers, json={
        "name": "状态流程测试", "protocol": {"seeds_per_dish": 5, "replicate_count": 1,
        "observation_period_days": 7, "sampling_rule": "first_germinated", "sample_count": 1,
        "sample_scope": "per_dish", "germination_criterion": "胚根可见"},
        "materials": [{"seed_lot_id": lot["id"]}], "dag_days": [0, 3],
    }).json()["experiment"]
    path = f"/api/experiments/{experiment['id']}"
    assert client.patch(path, headers=headers, json={"status": "ready"}).json()["status"] == "ready"
    assert client.patch(path, headers=headers, json={"status": "draft"}).json()["status"] == "draft"
    assert client.patch(path, headers=headers, json={"status": "ready"}).json()["status"] == "ready"
    assert client.patch(path, headers=headers, json={"status": "active"}).status_code == 409
    numbered = client.post(f"{path}/confirm-numbers", headers=headers)
    assert numbered.status_code == 200
    started = client.post(f"{path}/sowing/batch", headers=headers,
                          json={"sown_at": "2026-10-01T08:00:00+08:00",
                                "dish_ids": [dish["id"] for dish in numbered.json()["dishes"]]})
    assert started.status_code == 200, started.text
    assert client.get(path).json()["status"] == "active"


def test_ready_action_explains_incomplete_configuration(auth_client):
    client, headers = auth_client
    experiment = client.post("/api/experiments", headers=headers, json={"name": "尚未配置的实验"}).json()
    response = client.patch(f"/api/experiments/{experiment['id']}", headers=headers,
                            json={"status": "ready"})
    assert response.status_code == 422
    message = response.json()["detail"]
    assert "默认实验方案" in message
    assert "实验材料" in message
    assert "发芽后测定时间" in message
    assert client.get(f"/api/experiments/{experiment['id']}").json()["status"] == "draft"
