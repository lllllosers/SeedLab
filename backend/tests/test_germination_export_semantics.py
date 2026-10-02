from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from test_germination_execution import configured, start, batch, MORNING
from test_seedling_measurement import setup_experiment, observe


def export(client, headers, base):
    response = client.post('/api/export/experiments/workbook.xlsx', headers=headers,
                           json={'experiment_ids': [base.rsplit('/', 1)[-1]]})
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content), data_only=True)


def test_rate_blank_then_explicit_zero_then_last_observation_deleted(auth_client):
    client, headers = auth_client
    base, _, _ = configured(client, headers, replicates=1)
    dish = start(client, headers, base).json()['dishes'][0]
    book = export(client, headers, base)
    assert list(book['02_发芽率汇总'].values)[1][9:11] == (None, None)
    book.close()
    observed = batch(client, headers, base, MORNING, [{'dish_id': dish['id'], 'new_germinated_count': 0}]).json()
    book = export(client, headers, base)
    assert list(book['02_发芽率汇总'].values)[1][9:11] == (0, 0)
    book.close()
    observation_id = observed['created'][0]['id']
    assert client.delete(base + '/observations/' + observation_id, headers=headers).status_code == 204
    summary = client.get(base + '/execution').json()
    assert summary['cumulative_germinated'] is None and summary['germination_rate'] is None
    assert all(summary['dishes'][0][key] is None for key in (
        'cumulative_germinated', 'remaining_ungerminated', 'germination_rate'))
    book = export(client, headers, base)
    assert list(book['02_发芽率汇总'].values)[1][9:11] == (None, None)
    book.close()


def test_wide_keeps_all_unmeasured_samples_and_preserves_zero(auth_client):
    client, headers = auth_client
    base, dish, _, _ = setup_experiment(client, headers, days=(3, 7, 14))
    sample = observe(client, headers, base, dish, datetime.now(timezone.utc) - timedelta(days=3))
    book = export(client, headers, base)
    assert book['04_幼苗测定长表'].max_row == 1
    rows = list(book['05_幼苗测定宽表'].values)
    assert len(rows) == 2
    assert rows[1][3:5] == ('001', '001-01')
    assert rows[1][8:] == (None,) * 6
    book.close()
    task = next(t for t in client.get(base + '/measurement-tasks').json()['tasks']
                if t['day_after_germination'] == 3)
    response = client.post(base + '/measurements', headers=headers, json={
        'sample_id': sample['id'], 'timepoint_id': task['timepoint_id'],
        'root_length_mm': 0, 'shoot_length_mm': 0, 'measured_at': datetime.now(timezone.utc).isoformat()})
    assert response.status_code == 201, response.text
    book = export(client, headers, base)
    assert book['04_幼苗测定长表'].max_row == 2
    row = list(book['05_幼苗测定宽表'].values)[1]
    assert row[8:] == (0, 0, None, None, None, None)
    assert book['05_幼苗测定宽表']['I2'].data_type == 'n'
    assert book['05_幼苗测定宽表']['J2'].data_type == 'n'
    book.close()
