"""Small constructed designs verify the population independently of legacy data."""
from collections import Counter
from datetime import date, datetime, timedelta

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest
from sqlalchemy import func, inspect, select
from sqlalchemy.orm import Session

from app.db.base import Base
from app.db.session import make_engine
from app.models import (Experiment, ExperimentMaterial, ExperimentProtocol, GerminationDish,
                        MeasurementTimepoint, SeedlingMeasurement, SeedlingSample, SeedLot, Taxon)
from app.services.measurement_slots import build_measurement_slots
from app.core.config import get_settings
from app.services.local_time import utc_naive
from app.services.measurement_query import slots
from app.services.seedling_measurement import scheduled_date
from app.services.workbook_export import build


@pytest.fixture
def design(tmp_path):
    engine = make_engine(f"sqlite:///{(tmp_path / 'slots.db').as_posix()}")
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as db:
        taxon = Taxon(code='SP-0001', common_name='测试草', scientific_name='Test species')
        db.add(taxon); db.flush()
        lot = SeedLot(code='LOT-2026-001', taxon_id=taxon.id)
        db.add(lot); db.flush()
        experiment = Experiment(code='GER-202608-001', name='计划槽位测试', status='active',
                                planned_start_date=date(2026, 8, 3))
        db.add(experiment); db.flush()
        protocol = ExperimentProtocol(experiment_id=experiment.id, seeds_per_dish=50, replicate_count=1,
            sampling_rule='first_germinated', sample_count=10, sample_scope='per_dish')
        material = ExperimentMaterial(experiment_id=experiment.id, seed_lot_id=lot.id, experiment_number=1)
        db.add_all([protocol, material]); db.flush()
        dish = GerminationDish(material_id=material.id, code='INTERNAL-M001-R01', replicate_no=1,
                               label='R1', seed_count=50, sown_at=datetime(2026, 8, 3))
        points = [MeasurementTimepoint(experiment_id=experiment.id, day_after_germination=day) for day in (3, 7, 14)]
        db.add_all([dish, *points]); db.commit()
        yield db, experiment, protocol, material, dish, points
    engine.dispose()


def counts(db):
    return {table.name: db.scalar(select(func.count()).select_from(table)) for table in Base.metadata.sorted_tables}


def table(book, name):
    rows = list(book[name].values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def obtained(db, dish, number, when=None):
    sample = SeedlingSample(dish_id=dish.id, sample_number=number,
                            germinated_at=when or datetime(2026, 8, 8))
    db.add(sample); db.flush()
    return sample


@pytest.mark.parametrize('number_obtained', [0, 5, 10])
def test_zero_partial_and_full_designs_keep_all_slots_without_writing(design, number_obtained):
    db, experiment, _, _, dish, _ = design
    for number in range(1, number_obtained + 1):
        obtained(db, dish, number)
    db.commit()
    before = counts(db)
    canonical = build_measurement_slots(db, [experiment.id])
    assert not db.new and not db.dirty and not db.deleted
    assert all(inspect(slot, raiseerr=False) is None for slot in canonical.seedling_slots)
    assert all(inspect(slot, raiseerr=False) is None for slot in canonical.measurement_slots)
    assert len(canonical.seedling_slots) == 10 and len(canonical.measurement_slots) == 30
    assert sum(s.sample is not None for s in canonical.seedling_slots) == number_obtained
    assert Counter(s.data_status for s in canonical.measurement_slots) == {
        **({'无测定记录': number_obtained * 3} if number_obtained else {}),
        **({'无实际幼苗': (10-number_obtained) * 3} if number_obtained < 10 else {})}
    book = load_workbook(build(db, [experiment.id]), data_only=True)
    wide = table(book, '05_幼苗测定宽表')
    long = table(book, '04_幼苗测定长表')
    assert len(wide) == 10 and len(long) == 30
    assert [row['幼苗编号'] for row in wide] == [f'001-{n:02d}' for n in range(1, 11)]
    assert [row['是否已有实际幼苗'] for row in wide] == ['是'] * number_obtained + ['否'] * (10-number_obtained)
    for row in wide:
        assert (row['发芽判定时间'] is not None) == (row['是否已有实际幼苗'] == '是')
        assert all(row[column] is None for column in ('RL3', 'SL3', 'RL7', 'SL7', 'RL14', 'SL14'))
    assert len(table(book, '02_发芽率汇总')) == 1
    assert table(book, '02_发芽率汇总')[0]['累计发芽数'] is None
    assert table(book, '02_发芽率汇总')[0]['发芽率（%）'] is None
    book.close()
    assert counts(db) == before
    assert before['seedling_samples'] == number_obtained and before['seedling_measurements'] == 0


def test_partial_dag_all_dag_missing_and_numeric_zero(design):
    db, experiment, _, _, dish, points = design
    first = obtained(db, dish, 1)
    obtained(db, dish, 2)  # A real seedling with no measurements at any DAG.
    for point, root, shoot in ((points[0], 0, 7), (points[2], 3, 0)):
        db.add(SeedlingMeasurement(sample_id=first.id, timepoint_id=point.id,
            root_length_mm=root, shoot_length_mm=shoot, measured_at=first.germinated_at+timedelta(days=point.day_after_germination)))
    db.commit()
    book = load_workbook(build(db, [experiment.id]), data_only=True)
    long = table(book, '04_幼苗测定长表')
    assert Counter(row['数据状态'] for row in long) == {'已测定': 2, '无测定记录': 4, '无实际幼苗': 24}
    missing = next(row for row in long if row['幼苗编号'] == '001-01' and row['DAG'] == 7)
    assert missing['数据状态'] == '无测定记录' and missing['根长（mm）'] is None and missing['苗长（mm）'] is None
    assert missing['实际测定时间'] is None and missing['延迟天数'] is None
    assert missing['发芽判定时间'] is not None
    unmeasured = [row for row in long if row['幼苗编号'] == '001-02']
    assert len(unmeasured) == 3 and all(row['数据状态'] == '无测定记录' for row in unmeasured)
    wide = table(book, '05_幼苗测定宽表')
    assert (wide[0]['RL3'], wide[0]['SL3'], wide[0]['RL7'], wide[0]['SL7'], wide[0]['RL14'], wide[0]['SL14']) == (0, 7, None, None, 3, 0)
    assert book['05_幼苗测定宽表']['I2'].data_type == 'n'
    assert book['05_幼苗测定宽表']['N2'].data_type == 'n'
    for row in unmeasured:
        assert row['根长（mm）'] is None and row['苗长（mm）'] is None
    assert wide[1]['是否已有实际幼苗'] == '是' and wide[2]['是否已有实际幼苗'] == '否'
    book.close()


def test_per_dish_overrides_pending_and_cancelled_replicates_are_in_design(design):
    db, experiment, protocol, material, dish, _ = design
    protocol.sample_count = 4; protocol.replicate_count = 2
    material.sample_count_override = 3; material.replicate_count_override = 3
    dish.sown_at = None; dish.cancelled_at = datetime(2026, 8, 4)
    second = GerminationDish(material_id=material.id, code='INTERNAL-M001-R02', replicate_no=2,
                             label='R2', seed_count=50, sown_at=datetime(2026, 8, 3))
    db.add(second); db.flush()
    sample = obtained(db, second, 2)  # Preserve the real sequence gap.
    db.commit()
    before = counts(db)
    result = build_measurement_slots(db, [experiment.id])
    assert len(result.seedling_slots) == 9 and len(result.measurement_slots) == 27
    actual = next(s for s in result.seedling_slots if s.sample)
    assert actual.sample.id == sample.id and actual.seedling_number == '001-2-02'
    assert any(s.dish is None and s.dish_number == '001-3' for s in result.seedling_slots)
    assert len([s for s in result.seedling_slots if s.replicate_no == 1]) == 3
    book = load_workbook(build(db, [experiment.id]), data_only=True)
    assert len(table(book, '05_幼苗测定宽表')) == 9
    assert len(table(book, '04_幼苗测定长表')) == 27
    assert 'INTERNAL-' not in str([list(sheet.values) for sheet in book])
    book.close()
    assert counts(db) == before and before['germination_dishes'] == 2


def test_per_material_pool_does_not_multiply_target_or_invent_replicates(design):
    db, experiment, protocol, material, dish, _ = design
    protocol.sample_scope = 'per_material'; protocol.sample_count = 4; protocol.replicate_count = 2
    second = GerminationDish(material_id=material.id, code='INTERNAL-M001-R02', replicate_no=2,
                             label='R2', seed_count=50, sown_at=datetime(2026, 8, 3))
    db.add(second); db.flush()
    later = obtained(db, dish, 1, datetime(2026, 8, 9))
    earlier = obtained(db, second, 1, datetime(2026, 8, 8))
    db.commit()
    before = counts(db)
    result = build_measurement_slots(db, [experiment.id])
    assert len(result.seedling_slots) == 4 and len(result.measurement_slots) == 12
    assert [s.sample.id if s.sample else None for s in result.seedling_slots] == [earlier.id, later.id, None, None]
    assert [s.seedling_number for s in result.seedling_slots] == ['001-2-01', '001-1-01', None, None]
    book = load_workbook(build(db, [experiment.id]), data_only=True)
    wide = table(book, '05_幼苗测定宽表')
    assert [row['计划取样序号'] for row in wide] == [1, 2, 3, 4]
    assert all(row['取样范围'] == '每材料' for row in wide)
    assert all(row['培养皿现场编号'] is None and row['幼苗编号'] is None for row in wide[2:])
    assert len(table(book, '04_幼苗测定长表')) == 12
    book.close()
    assert counts(db) == before


def test_multi_experiment_dags_do_not_expand_another_experiments_design(design):
    db, first, _, material, _, _ = design
    second = Experiment(code='GER-202608-002', name='另一测定方案')
    db.add(second); db.flush()
    db.add_all([ExperimentProtocol(experiment_id=second.id, seeds_per_dish=50, replicate_count=1,
        sampling_rule='first_germinated', sample_count=2, sample_scope='per_dish'),
        ExperimentMaterial(experiment_id=second.id, seed_lot_id=material.seed_lot_id, experiment_number=1),
        MeasurementTimepoint(experiment_id=second.id, day_after_germination=5)])
    db.commit()
    before = counts(db)
    book = load_workbook(build(db, [first.id, second.id]), data_only=True)
    wide, long = table(book, '05_幼苗测定宽表'), table(book, '04_幼苗测定长表')
    assert len(wide) == 12 and len(long) == 32
    assert len(table(book, '02_发芽率汇总')) == 2
    assert all(row['DAG'] == 5 for row in long if second.code in row['来源实验'])
    assert all(row['DAG'] in (3, 7, 14) for row in long if first.code in row['来源实验'])
    assert {'RL3', 'RL5', 'RL7', 'RL14'} <= wide[0].keys()
    book.close()
    assert counts(db) == before


@pytest.mark.parametrize('scope', ['per_dish', 'per_material'])
def test_inconsistent_sampling_design_cannot_silently_drop_actual_seedlings(design, scope):
    db, experiment, protocol, _, dish, _ = design
    protocol.sample_scope = scope
    for number in range(1, 12):
        obtained(db, dish, number)
    db.commit()
    before = counts(db)
    with pytest.raises(HTTPException, match='避免遗漏实际幼苗'):
        build_measurement_slots(db, [experiment.id])
    assert counts(db) == before


def test_incomplete_design_has_actionable_error_instead_of_disappearing_material(design):
    db, experiment, protocol, _, _, _ = design
    protocol.sample_count = None
    db.commit()
    with pytest.raises(HTTPException, match='完善实验方案'):
        build(db, [experiment.id])


def test_unconfirmed_numbers_still_have_distinguishable_planned_replicates(design):
    db, experiment, protocol, material, _, _ = design
    protocol.replicate_count = 2
    material.experiment_number = None
    db.commit()
    before = counts(db)
    book = load_workbook(build(db, [experiment.id]), data_only=True)
    wide = table(book, '05_幼苗测定宽表')
    assert len(wide) == 20
    assert Counter(row['计划培养皿重复'] for row in wide) == {1: 10, 2: 10}
    assert all(row['培养皿现场编号'] is None and row['幼苗编号'] is None for row in wide)
    assert len({(row['汇总编号'], row['计划培养皿重复'], row['计划取样序号']) for row in wide}) == 20
    book.close()
    assert counts(db) == before


def test_selected_empty_experiment_keeps_configured_dag_header_without_inventing_material(design):
    db, first, _, _, _, _ = design
    empty = Experiment(code='GER-202608-002', name='尚未添加材料的实验')
    db.add(empty); db.flush()
    db.add(MeasurementTimepoint(experiment_id=empty.id, day_after_germination=21))
    db.commit()
    before = counts(db)
    book = load_workbook(build(db, [first.id, empty.id]), data_only=True)
    wide = table(book, '05_幼苗测定宽表')
    assert len(wide) == 10 and 'RL21' in wide[0] and 'SL21' in wide[0]
    long = table(book, '04_幼苗测定长表')
    assert len(long) == 30 and all(row['DAG'] in (3, 7, 14) for row in long)
    assert all(row['RL21'] is None and row['SL21'] is None for row in wide)
    book.close()
    assert counts(db) == before


@pytest.mark.parametrize('zone, timestamp, expected_day', [
    ('Asia/Shanghai', '2026-01-31T16:20:00+00:00', '2026-02-01'),
    ('Asia/Shanghai', '2026-12-31T15:59:00+00:00', '2026-12-31'),
    ('Asia/Shanghai', '2026-12-31T16:00:00+00:00', '2027-01-01'),
    ('America/New_York', '2026-03-08T06:30:00+00:00', '2026-03-08'),
    ('America/New_York', '2026-11-01T05:30:00+00:00', '2026-11-01'),
    ('UTC', '2026-02-28T23:59:00', '2026-02-28'),
    ('Asia/Shanghai', None, None),
])
def test_python_and_sql_dag_calendar_parity(design, monkeypatch, zone, timestamp, expected_day):
    db, experiment, _, _, dish, _ = design
    monkeypatch.setenv('SEEDLAB_TIMEZONE', zone)
    get_settings.cache_clear()
    try:
        when = utc_naive(datetime.fromisoformat(timestamp)) if timestamp else None
        sample = SeedlingSample(dish_id=dish.id, sample_number=1, germinated_at=when)
        point = MeasurementTimepoint(experiment_id=experiment.id, day_after_germination=0)
        db.add_all([sample, point])
        db.commit()
        projection = slots().c
        rows = db.execute(select(projection).where(projection.sample_id == sample.id)
                          .order_by(projection.day_after_germination)).mappings().all()
        assert [row['day_after_germination'] for row in rows] == [0, 3, 7, 14]
        for row in rows:
            dag = row['day_after_germination']
            expected = (date.fromisoformat(expected_day) + timedelta(days=dag)).isoformat() if expected_day else None
            python_day = scheduled_date(sample.germinated_at, dag)
            assert row['scheduled_date'] == (python_day.isoformat() if python_day else None) == expected
        if when is None:
            assert all(row['status'] == 'unschedulable' for row in rows)
    finally:
        get_settings.cache_clear()


def test_workbook_schema_and_independent_zero_blank_unavailable_states(design):
    db, experiment, _, _, dish, points = design
    for number, root, shoot, root_na, shoot_na in (
        (1, 0, None, False, True), (2, None, 0, True, False), (3, None, None, True, True)):
        sample = obtained(db, dish, number)
        db.add(SeedlingMeasurement(sample_id=sample.id, timepoint_id=points[0].id,
            root_length_mm=root, shoot_length_mm=shoot, root_unavailable=root_na,
            shoot_unavailable=shoot_na, measured_at=sample.germinated_at + timedelta(days=3)))
    obtained(db, dish, 4)  # Actual sample without a measurement is a separate fact.
    db.commit()
    before = counts(db)
    canonical = build_measurement_slots(db, [experiment.id])
    assert Counter(stage.data_status for stage in canonical.measurement_slots) == {
        '已测定': 3, '无测定记录': 9, '无实际幼苗': 18}
    expected_headers = {
        '01_材料总表': ('汇总编号', '来源实验', '原实验编号', '中文名', '学名', '科', '属', '生活型',
            '原始材料编号', '系统物种编号', '系统种子批次编号', '来源', '采集/获得日期', '数量', '备注'),
        '02_发芽率汇总': ('汇总编号', '来源实验', '原实验编号', '中文名', '学名', '原始材料编号',
            '种子批次', '实际已置床培养皿数', '实际置床种子数', '累计发芽数', '发芽率（%）'),
        '03_发芽原始记录': ('汇总编号', '来源实验', '原实验编号', '培养皿现场编号', '中文名', '学名',
            '实际置床时间', '巡检时间', '本次新增发芽数', '累计发芽数', '发芽率'),
        '04_幼苗测定长表': ('汇总编号', '来源实验', '原实验编号', '培养皿现场编号', '幼苗编号', '位置标签',
            '中文名', '学名', '发芽判定时间', 'DAG', '计划测定日期', '实际测定时间', '延迟天数',
            '根长（mm）', '根长状态', '苗长（mm）', '苗长状态', '备注', '数据状态', '计划取样序号', '取样范围', '计划培养皿重复'),
        '05_幼苗测定宽表': ('汇总编号', '来源实验', '原实验编号', '培养皿现场编号', '幼苗编号', '位置标签',
            '中文名', '学名', 'RL3', 'SL3', 'RL7', 'SL7', 'RL14', 'SL14', '发芽判定时间',
            '是否已有实际幼苗', '计划取样序号', '取样范围', '计划培养皿重复'),
        '06_导出说明': ('项目', '说明'),
    }
    book = load_workbook(build(db, [experiment.id]), data_only=True)
    try:
        assert book.sheetnames == list(expected_headers)
        for name, header in expected_headers.items():
            assert tuple(cell.value for cell in book[name][1]) == header
        wide = table(book, '05_幼苗测定宽表')
        assert [(row['RL3'], row['SL3']) for row in wide[:5]] == [
            (0, 'NA'), ('NA', 0), ('NA', 'NA'), (None, None), (None, None)]
        assert [row['是否已有实际幼苗'] for row in wide[:5]] == ['是', '是', '是', '是', '否']
        long = [row for row in table(book, '04_幼苗测定长表') if row['DAG'] == 3]
        assert [(row['根长状态'], row['苗长状态'], row['数据状态']) for row in long[:5]] == [
            ('已测', '无法测量', '已测定'), ('无法测量', '已测', '已测定'),
            ('无法测量', '无法测量', '已测定'), ('无测定记录', '无测定记录', '无测定记录'),
            ('无实际幼苗', '无实际幼苗', '无实际幼苗')]
        assert book['05_幼苗测定宽表']['I2'].data_type == 'n'
        assert book['05_幼苗测定宽表']['J2'].data_type == 's'
        assert book['05_幼苗测定宽表']['J3'].data_type == 'n'
        assert all(row['RL7'] is None and row['SL7'] is None for row in wide)
    finally:
        book.close()
    assert counts(db) == before
