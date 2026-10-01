"""Round 2: query feedback, material selection, import handoff and user catalog."""
from datetime import datetime, timezone
from io import BytesIO
import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from app.models import Experiment, SeedLot, Taxon
from test_measurement_workflow_and_lifecycle import fixture_engine, populate
from test_seedling_measurement import payload
from test_material_import_and_execution import material_file, row, confirm, config, ready
from app.services.catalog_export import TAXON_HEADERS, LOT_HEADERS


@pytest.mark.parametrize('state', ['draft', 'ready', 'active', 'completed', 'cancelled'])
def test_unparticipated_excludes_every_existing_experiment_state(auth_client, tmp_path, state):
    client, headers = auth_client
    taxon = client.post('/api/taxa', json={'scientific_name': 'Aster test'}, headers=headers).json()
    unused = client.post('/api/seed-lots', json={'taxon_id': taxon['id']}, headers=headers).json()
    used = client.post('/api/seed-lots', json={'taxon_id': taxon['id']}, headers=headers).json()
    exp = config(client, headers, [used['id']])['experiment']
    engine = fixture_engine(tmp_path)
    with Session(engine) as db:
        db.get(Experiment, exp['id']).status = state
        db.commit()
    engine.dispose()
    all_lots = client.get('/api/experiments/available-seed-lots').json()
    result = client.get('/api/experiments/available-seed-lots', params={'unparticipated': True}).json()
    assert {item['id'] for item in all_lots} == {used['id'], unused['id']}
    assert [item['id'] for item in result] == [unused['id']]
    assert result[0]['sort_rank'] == next(item['sort_rank'] for item in all_lots if item['id'] == unused['id'])
    if state in {'draft', 'ready'}:
        assert client.delete('/api/experiments/' + exp['id'], headers=headers).status_code == 204
        assert len(client.get('/api/experiments/available-seed-lots', params={'unparticipated': True}).json()) == 2


def test_selected_preview_order_matches_final_numbering_across_filters(auth_client):
    client, headers = auth_client
    for n, name in enumerate(['稗', '阿尔泰狗娃花', '巴天酸模']):
        taxon = client.post('/api/taxa', json={'scientific_name': f'Species {n}', 'common_name': name}, headers=headers).json()
        client.post('/api/seed-lots', json={'taxon_id': taxon['id'], 'source_code': '000' + str(n)}, headers=headers)
    available = client.get('/api/experiments/available-seed-lots').json()
    assert [lot['taxon_common_name'] for lot in available] == ['阿尔泰狗娃花', '巴天酸模', '稗']
    for lot in available:
        filtered = client.get('/api/experiments/available-seed-lots', params={'q': lot['taxon_common_name']}).json()
        assert filtered[0]['sort_rank'] == lot['sort_rank']
    exp = config(client, headers, [lot['id'] for lot in reversed(available)])['experiment']
    ready(client, headers, exp['id'])
    numbered = client.get('/api/experiments/' + exp['id'] + '/configuration').json()
    assert [(m['seed_lot_id'], m['experiment_number']) for m in numbered['materials']] == [(lot['id'], n + 1) for n, lot in enumerate(available)]


def test_import_handoff_all_created_updated_and_deduplicated(auth_client):
    client, headers = auth_client
    taxon = client.post('/api/taxa', json={'scientific_name': 'Setaria viridis', 'common_name': '狗尾草'}, headers=headers).json()
    existing = client.post('/api/seed-lots', json={'taxon_id': taxon['id'], 'source_code': '0001'}, headers=headers).json()
    content = material_file([row('Setaria viridis', '0001', source='武威'),
        row('Setaria viridis', '0002', quantity=0), row('Poa annua', '0003'), row('Poa annua', '0003')])
    response = confirm(client, headers, content, {'5': 'reuse-row:4'})
    assert response.status_code == 200, response.text
    result = response.json()
    assert len(result['all_seed_lot_ids']) == result['total_material_count'] == 3
    assert len(result['created_seed_lot_ids']) == result['created_material_count'] == 2
    assert existing['id'] in result['all_seed_lot_ids'] and existing['id'] not in result['created_seed_lot_ids']
    assert result['existing_material_count'] == 1 and result['updated_material_count'] == 1
    lots = client.get('/api/seed-lots').json()
    assert next(lot for lot in lots if lot['source_code'] == '0002')['id'] in result['created_seed_lot_ids']
    assert next(lot for lot in lots if lot['source_code'] == '0002')['quantity'] == 0
    # Same logical content in a distinct file, with no new batch.
    no_new = confirm(client, headers, material_file([row('Setaria viridis', '0001', source='武威', quantity=None)]))
    assert no_new.status_code == 200, no_new.text
    assert no_new.json()['created_seed_lot_ids'] == [] and no_new.json()['created_material_count'] == 0
    assert no_new.json()['all_seed_lot_ids'] == [existing['id']]


def test_integrated_import_200_rows_and_invalid_sheet_is_atomic(auth_client, tmp_path):
    client, headers = auth_client
    content = material_file([row(f'Species {n}', f'{n:04d}', f'物种{n}', quantity=0 if n == 0 else None) for n in range(200)])
    result = confirm(client, headers, content)
    assert result.status_code == 200, result.text
    assert result.json()['total_material_count'] == result.json()['created_material_count'] == 200
    engine = fixture_engine(tmp_path)
    with engine.connect() as conn:
        assert conn.exec_driver_sql("SELECT COUNT(*) FROM audit_logs WHERE entity_type='SeedLot' AND action='import'").scalar() == 200
    engine.dispose()
    bad = material_file([row('New valid species', 'valid'), row('New invalid species', 'invalid', quantity=-1)])
    report = client.post('/api/import/materials/preview', headers=headers, files={'file': ('bad.xlsx', bad)}).json()
    assert report['stats']['errors'] == 1 and '第 3 行' in ' '.join(report['rows'][1]['errors'])
    assert confirm(client, headers, bad).status_code == 422
    assert len(client.get('/api/seed-lots').json()) == 200
    assert len(client.get('/api/taxa').json()) == 200


def test_record_feedback_counts_filtered_materials_across_pages(auth_client, tmp_path):
    client, headers = auth_client
    base, ids = populate(tmp_path, 3, 2)
    for task in client.get(base + '/measurement-tasks').json()['tasks']:
        if task['day_after_germination'] == 0:
            assert client.post(base + '/measurements', json=payload(task, datetime.now(timezone.utc)), headers=headers).status_code == 201
    result = client.get(base + '/measurement-records', params={'page_size': 2}).json()
    assert result['total'] == 6 and result['material_count'] == 3 and len(result['items']) == 2
    for params, total, count in [({'q': '001'}, 2, 1), ({'material_ids': ids[:2]}, 4, 2), ({'dag': 7}, 0, 0), ({}, 6, 3)]:
        response = client.get(base + '/measurement-records', params=params)
        assert response.status_code == 200
        assert (response.json()['total'], response.json()['material_count']) == (total, count)


def test_catalog_ledger_chinese_headers_text_ids_zero_unicode_and_status(auth_client):
    client, headers = auth_client
    taxon = client.post('/api/taxa', json={'scientific_name': 'Setaria viridis', 'common_name': '狗尾草', 'family': '禾本科', 'genus': '狗尾草属', 'life_form': '一年生', 'notes': '中文备注'}, headers=headers).json()
    lot = client.post('/api/seed-lots', json={'taxon_id': taxon['id'], 'source_code': '000123', 'quantity': 0, 'source': '武威', 'notes': '=原始文本'}, headers=headers).json()
    client.patch('/api/taxa/' + taxon['id'], json={'is_active': False}, headers=headers)
    result = client.get('/api/export/catalog.xlsx')
    assert result.status_code == 200
    workbook = load_workbook(BytesIO(result.content))
    assert workbook.sheetnames == ['01_物种', '02_种子批次']
    assert tuple(cell.value for cell in workbook.worksheets[0][1]) == TAXON_HEADERS
    assert tuple(cell.value for cell in workbook.worksheets[1][1]) == LOT_HEADERS
    assert workbook.worksheets[0]['G2'].value == '已停用'
    sheet = workbook.worksheets[1]
    assert sheet['E2'].value == '000123' and sheet['E2'].data_type == 's' and sheet['E2'].number_format == '@'
    assert sheet['H2'].value == 0 and sheet['C2'].value == '狗尾草' and sheet['I2'].value == '使用中'
    assert sheet['J2'].value == '=原始文本' and sheet['J2'].data_type == 's'
    assert sheet['A2'].value == lot['code'] and sheet['B2'].value == taxon['code']
    workbook.close()
    for path in ('/api/export/taxa.csv', '/api/export/taxa-template.xlsx', '/api/export/seed-lots-template.xlsx'):
        assert client.get(path).status_code == 404
    for path in ('/api/import/taxa', '/api/import/seed-lots'):
        assert client.post(path, headers=headers).status_code == 404
