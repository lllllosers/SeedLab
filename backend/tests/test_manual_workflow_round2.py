import json
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from alembic import command
from alembic.config import Config
from pathlib import Path

from app.core.config import get_settings

from app.db.session import make_engine
from app.services.material_import import COLUMNS


def material_file(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def row(scientific, code, common=None, *, taxon_code=None, source=None, collected=None,
        quantity=50, family=None, genus=None, life_form=None):
    return (taxon_code, common, scientific, family, genus, life_form, code, source,
            collected, quantity, None, None)


def preview(client, headers, content, mode="complete"):
    response = client.post("/api/import/materials/preview", headers=headers,
                           data={"mode": mode},
                           files={"file": ("materials.xlsx", content,
                                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200, response.text
    return response.json()


def confirm(client, headers, content, decisions=None, mode="complete"):
    return client.post("/api/import/materials/confirm", headers=headers,
                       data={"decisions": json.dumps(decisions or {}), "mode": mode},
                       files={"file": ("materials.xlsx", content,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def config(client, headers, lot_ids, name="第二轮实验", replicates=2):
    response = client.post("/api/experiments/configured", headers=headers, json={
        "name": name,
        "protocol": {"seeds_per_dish": 20, "replicate_count": replicates,
                     "observation_period_days": 7, "sampling_rule": "first_germinated",
                     "sample_count": 2, "sample_scope": "per_dish",
                     "germination_criterion": "胚根可见"},
        "materials": [{"seed_lot_id": item} for item in lot_ids], "dag_days": [0, 5],
    })
    assert response.status_code == 201, response.text
    return response.json()


def ready(client, headers, experiment_id):
    assert client.patch(f"/api/experiments/{experiment_id}", headers=headers,
                        json={"status": "ready"}).status_code == 200
    response = client.post(f"/api/experiments/{experiment_id}/confirm-numbers", headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def test_complete_list_incremental_import_and_hash(auth_client, tmp_path):
    client, headers = auth_client
    template = load_workbook(BytesIO(client.get("/api/export/materials-template.xlsx").content), read_only=True)
    assert tuple(next(template.active.values)) == COLUMNS
    template.close()
    first_rows = [row(f"Species {index}", f"A{index:03d}", f"物种{index}") for index in range(10)]
    first = material_file(first_rows)
    assert preview(client, headers, first)["stats"]["new"] == 10
    imported = confirm(client, headers, first)
    assert imported.status_code == 200, imported.text
    assert imported.json()["new_taxa"] == 10
    assert imported.json()["new_lots"] == 10
    assert preview(client, headers, first)["already_imported"] is True
    assert confirm(client, headers, first).status_code == 409
    full = material_file(first_rows + [row(f"Species {index}", f"A{index:03d}", f"物种{index}") for index in range(10, 30)])
    report = preview(client, headers, full)
    assert report["stats"]["registered"] == 10
    assert report["stats"]["new"] == 20
    imported = confirm(client, headers, full)
    assert imported.status_code == 200, imported.text
    assert imported.json()["new_lots"] == 20
    assert len(imported.json()["selected_lot_ids"]) == 20
    assert len(client.get("/api/seed-lots").json()) == 30
    engine = make_engine(f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    with engine.connect() as connection:
        assert connection.execute(text("SELECT COUNT(*) FROM import_jobs WHERE file_hash IS NOT NULL")).scalar() == 2
        assert connection.execute(text("SELECT COUNT(*) FROM audit_logs WHERE entity_type='SeedLot' AND action='import'")).scalar() == 30
    engine.dispose()


def test_import_matching_fills_and_conflicts(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers, json={
        "scientific_name": "Setaria viridis", "common_name": "狗尾草", "life_form": "一年生"}).json()
    lot = client.post("/api/seed-lots", headers=headers, json={
        "taxon_id": taxon["id"], "source_code": "A027", "quantity": 50}).json()
    content = material_file([row(" setaria VIRIDIS ", "A027", "狗尾草", genus="狗尾草属",
                                 life_form="多年生", quantity=80)])
    report = preview(client, headers, content)
    item = report["rows"][0]
    assert item["status"] == "updatable"
    assert "属" in item["fills"]
    assert "生活型" in item["conflicts"]
    assert "数量（粒）" in item["conflicts"]
    response = confirm(client, headers, content)
    assert response.status_code == 200, response.text
    assert response.json()["new_lots"] == 0
    assert client.get(f"/api/taxa/{taxon['id']}").json()["genus"] == "狗尾草属"
    assert client.get(f"/api/taxa/{taxon['id']}").json()["life_form"] == "一年生"
    assert client.get(f"/api/seed-lots/{lot['id']}").json()["quantity"] == 50
    wrong = material_file([row("Setaria viridis", "B001", taxon_code="SP-9999")])
    assert "不存在" in preview(client, headers, wrong)["rows"][0]["errors"][0]
    assert confirm(client, headers, wrong).status_code == 422
    same_chinese = material_file([row("Setaria faberi", "C001", "狗尾草")])
    assert preview(client, headers, same_chinese)["stats"]["new_taxa"] == 1
    assert confirm(client, headers, same_chinese).json()["new_taxa"] == 1


def test_import_ambiguous_and_duplicate_rows_need_decision(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers,
                        json={"scientific_name": "Setaria viridis"}).json()
    lot = client.post("/api/seed-lots", headers=headers,
                      json={"taxon_id": taxon["id"], "source": "武威",
                            "collected_at": "2026-09-01T00:00:00Z"}).json()
    content = material_file([row("Setaria viridis", None, source="武威", collected="2026-09-01"),
                             row("Setaria viridis", None, source="武威", collected="2026-09-01")])
    report = preview(client, headers, content)
    assert [item["status"] for item in report["rows"]] == ["updatable", "confirm"]
    assert report["rows"][1]["candidate_row"] == 2
    assert confirm(client, headers, content).status_code == 422
    response = confirm(client, headers, content, {"2": "reuse", "3": "reuse"})
    assert response.status_code == 200, response.text
    assert response.json()["new_lots"] == 0
    assert lot["id"] in response.json()["selected_lot_ids"]


def test_import_multiple_possible_lots_requires_specific_choice(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers,
                        json={"scientific_name": "Setaria viridis"}).json()
    lots = [client.post("/api/seed-lots", headers=headers, json={
        "taxon_id": taxon["id"], "source": "武威",
        "collected_at": "2026-09-01T00:00:00Z"}).json() for _ in range(2)]
    content = material_file([row("Setaria viridis", None, source="武威", collected="2026-09-01")])
    report = preview(client, headers, content)
    assert report["rows"][0]["status"] == "confirm"
    assert {candidate["id"] for candidate in report["rows"][0]["candidate_lots"]} == {lot["id"] for lot in lots}
    assert confirm(client, headers, content, {"2": "reuse"}).status_code == 422
    assert confirm(client, headers, content, {"2": 12}).status_code == 422
    chosen = confirm(client, headers, content, {"2": f"reuse:{lots[1]['id']}"})
    assert chosen.status_code == 200, chosen.text
    assert chosen.json()["selected_lot_ids"] == [lots[1]["id"]]


def test_catalog_deletion_respects_experiment_history(auth_client):
    client, headers = auth_client
    unused = client.post("/api/taxa", headers=headers, json={"scientific_name": "Unused species"}).json()
    assert client.delete(f"/api/taxa/{unused['id']}", headers=headers).status_code == 204
    taxon = client.post("/api/taxa", headers=headers, json={"scientific_name": "Setaria viridis"}).json()
    unused_lot = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxon["id"]}).json()
    assert client.delete(f"/api/taxa/{taxon['id']}", headers=headers).status_code == 409
    assert client.delete(f"/api/seed-lots/{unused_lot['id']}", headers=headers).status_code == 204
    used_lot = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxon["id"]}).json()
    config(client, headers, [used_lot["id"]])
    blocked = client.delete(f"/api/seed-lots/{used_lot['id']}", headers=headers)
    assert blocked.status_code == 409
    assert "停用" in blocked.json()["detail"]
    assert client.patch(f"/api/seed-lots/{used_lot['id']}", headers=headers,
                        json={"is_active": False}).status_code == 200
    assert client.delete(f"/api/taxa/{taxon['id']}", headers=headers).status_code == 409


def test_import_modes_and_batch_decisions_without_source_codes(auth_client):
    client, headers = auth_client
    taxa = [client.post("/api/taxa", headers=headers, json={"scientific_name": f"Species {i}"}).json()
            for i in range(3)]
    first = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxa[0]["id"],
                        "source": "采集地甲", "quantity": 20}).json()
    historical = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxa[0]["id"],
                              "source": "历史来源"}).json()
    assert client.patch(f"/api/seed-lots/{historical['id']}", headers=headers,
                        json={"is_active": False}).status_code == 200
    client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxa[1]["id"], "source": "采集地乙"})
    client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxa[1]["id"], "source": "采集地丙"})
    client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxa[2]["id"], "source": "采集地丁"})
    content = material_file([row("Species 0", None, source="采集地甲", quantity=20),
                             row("Species 1", None, source="采集地乙"),
                             row("Species 2", None, source="新来源")])
    report = preview(client, headers, content)
    assert [item["status"] for item in report["rows"]] == ["registered", "confirm", "confirm"]
    assert report["rows"][0]["candidate_lot_id"] == first["id"]
    assert confirm(client, headers, content).status_code == 422
    decisions = {str(item["line"]): "new" for item in report["rows"] if item["status"] == "confirm"}
    imported = confirm(client, headers, content, decisions)
    assert imported.status_code == 200, imported.text
    assert imported.json()["new_lots"] == 2
    new_content = material_file([row("Species 0", None, source="另一来源", quantity=10)])
    assert preview(client, headers, new_content, mode="new")["rows"][0]["status"] == "new"
    assert confirm(client, headers, new_content, mode="new").json()["new_lots"] == 1
    dated = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxa[0]["id"],
                         "source": "有日期的材料", "collected_at": "2026-09-01T00:00:00Z"}).json()
    exact = material_file([row("Species 0", None, source="有日期的材料", collected="2026-09-01")])
    duplicate = preview(client, headers, exact, mode="new")["rows"][0]
    assert duplicate["status"] == "confirm"
    assert dated["id"] in {item["id"] for item in duplicate["candidate_lots"]}


def test_import_bad_cell_reports_row_and_blocks_entire_file(auth_client):
    client, headers = auth_client
    content = material_file([
        row("Setaria viridis", "A001", "狗尾草"),
        row("Trifolium repens", "B001", "白三叶", quantity=-2),
    ])
    report = preview(client, headers, content)
    assert report["stats"]["total"] == 2
    assert report["stats"]["errors"] == 1
    assert report["rows"][0]["status"] == "new"
    assert report["rows"][1]["status"] == "error"
    assert "第 3 行" in report["rows"][1]["errors"][0]
    assert confirm(client, headers, content).status_code == 422
    assert client.get("/api/seed-lots").json() == []


def test_import_500_rows_and_atomic_rollback(auth_client, tmp_path, monkeypatch):
    client, headers = auth_client
    content = material_file([row("Setaria viridis", f"A{index:03d}", "狗尾草", quantity=index)
                             for index in range(500)])
    assert preview(client, headers, content)["stats"]["new"] == 500
    from app.services import material_import
    original = material_import.flush_or_conflict
    calls = 0
    def fail_midway(db):
        nonlocal calls
        calls += 1
        if calls == 138:
            raise RuntimeError("injected import failure")
        return original(db)
    monkeypatch.setattr(material_import, "flush_or_conflict", fail_midway)
    with pytest.raises(RuntimeError):
        confirm(client, headers, content)
    monkeypatch.setattr(material_import, "flush_or_conflict", original)
    assert client.get("/api/seed-lots").json() == []
    response = confirm(client, headers, content)
    assert response.status_code == 200, response.text
    assert response.json()["new_lots"] == 500
    engine = make_engine(f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    with engine.connect() as connection:
        assert connection.execute(text("SELECT COUNT(*) FROM seed_lots")).scalar() == 500
        assert connection.execute(text("SELECT COUNT(*) FROM import_jobs WHERE file_hash IS NOT NULL")).scalar() == 1
    engine.dispose()


def test_numbering_sort_lock_reopen_and_multiple_experiments(auth_client, tmp_path):
    client, headers = auth_client
    taxa = [client.post("/api/taxa", headers=headers, json={"scientific_name": sci,
             "common_name": common}).json() for common, sci in
            [("狗尾草", "Setaria viridis"), ("白三叶", "Trifolium repens")]]
    lots = [client.post("/api/seed-lots", headers=headers,
            json={"taxon_id": taxon_id, "source_code": code}).json() for taxon_id, code in
            [(taxa[0]["id"], "B"), (taxa[0]["id"], "A"), (taxa[1]["id"], "C")]]
    created = config(client, headers, [lot["id"] for lot in lots])
    materials = created["materials"]
    assert [item["source_code"] for item in materials] == ["C", "A", "B"]
    assert [item["preview_number"] for item in materials] == [1, 2, 3]
    experiment_id = created["experiment"]["id"]
    summary = ready(client, headers, experiment_id)
    assert [item["experiment_number"] for item in summary["materials"]] == [1, 2, 3]
    assert [item["field_number"] for item in summary["dishes"]] == ["001-1", "001-2", "002-1", "002-2", "003-1", "003-2"]
    assert all(item["sown_at"] is None for item in summary["dishes"])
    sheet = load_workbook(BytesIO(client.get(f"/api/experiments/{experiment_id}/sowing-sheet.xlsx").content))
    assert [sheet.active.cell(row, 1).value for row in (2, 4, 6)] == ["001", "002", "003"]
    sheet.close()
    assert client.post(f"/api/experiments/{experiment_id}/reopen-design", headers=headers).status_code == 200
    assert client.get(f"/api/experiments/{experiment_id}/execution").json()["dish_count"] == 0
    assert [item["experiment_number"] for item in client.get(f"/api/experiments/{experiment_id}/configuration").json()["materials"]] == [None] * 3
    assert client.patch(f"/api/experiments/{experiment_id}", headers=headers, json={"status": "ready"}).status_code == 200
    ready(client, headers, experiment_id)
    other = config(client, headers, [lots[0]["id"]], name="另一次实验")
    second = ready(client, headers, other["experiment"]["id"])
    assert second["materials"][0]["experiment_number"] == 1
    cancelled_id = second["dishes"][1]["id"]
    cancelled = client.post(f"/api/experiments/{other['experiment']['id']}/sowing/{cancelled_id}/cancel",
                            headers=headers, json={"reason": "培养皿破损"})
    assert cancelled.status_code == 200
    assert cancelled.json()["dishes"][1]["field_number"] == "001-2"
    assert client.post(f"/api/experiments/{other['experiment']['id']}/reopen-design", headers=headers).status_code == 200
    assert client.patch(f"/api/experiments/{other['experiment']['id']}", headers=headers,
                        json={"status": "ready"}).status_code == 200
    second = ready(client, headers, other["experiment"]["id"])
    cancelled_id = second["dishes"][1]["id"]
    assert client.post(f"/api/experiments/{other['experiment']['id']}/sowing/{cancelled_id}/cancel",
                       headers=headers, json={"reason": "培养皿破损"}).status_code == 200
    active = client.post(f"/api/experiments/{other['experiment']['id']}/sowing/batch", headers=headers,
                         json={"sown_at": "2026-09-28T09:00:00Z", "dish_ids": [second["dishes"][0]["id"]]})
    assert active.status_code == 200
    assert active.json()["seed_count"] == 20
    blocked = client.post(f"/api/experiments/{other['experiment']['id']}/observations/batch",
                          headers=headers, json={"observed_at": "2026-09-29T09:00:00Z", "entries": [
                              {"dish_id": cancelled_id, "new_germinated_count": 0}]})
    assert blocked.status_code == 409
    engine = make_engine(f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    with engine.connect() as connection:
        assert connection.exec_driver_sql("PRAGMA foreign_key_check").all() == []
    engine.dispose()


def test_multiday_sowing_observation_and_export(auth_client, tmp_path):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers,
                        json={"scientific_name": "Setaria viridis", "common_name": "狗尾草"}).json()
    lot = client.post("/api/seed-lots", headers=headers,
                      json={"taxon_id": taxon["id"], "source_code": "A027"}).json()
    first = config(client, headers, [lot["id"]], name="第一次实验")
    first_id = first["experiment"]["id"]
    summary = ready(client, headers, first_id)
    dishes = summary["dishes"]
    assert client.post(f"/api/experiments/{first_id}/observations/batch", headers=headers,
                       json={"observed_at": "2026-09-29T10:00:00Z", "entries": [
                           {"dish_id": dishes[0]["id"], "new_germinated_count": 0}]}).status_code == 409
    first_sow = client.post(f"/api/experiments/{first_id}/sowing/batch", headers=headers,
                            json={"sown_at": "2026-09-28T09:00:00Z", "dish_ids": [dishes[1]["id"]]})
    assert first_sow.status_code == 200, first_sow.text
    assert first_sow.json()["experiment"]["status"] == "active"
    assert first_sow.json()["experiment"]["started_at"] == "2026-09-28T09:00:00Z"
    assert first_sow.json()["seed_count"] == 20
    assert client.post(f"/api/experiments/{first_id}/reopen-design", headers=headers).status_code == 409
    assert client.post(f"/api/experiments/{first_id}/observations/batch", headers=headers,
                       json={"observed_at": "2026-09-29T10:00:00Z", "entries": [
                           {"dish_id": dishes[0]["id"], "new_germinated_count": 0}]}).status_code == 422
    second_sow = client.post(f"/api/experiments/{first_id}/sowing/batch", headers=headers,
                             json={"sown_at": "2026-09-29T08:00:00Z", "dish_ids": [dishes[0]["id"]]})
    assert second_sow.status_code == 200, second_sow.text
    observed_at = "2026-09-29T10:00:00Z"
    observed = client.post(f"/api/experiments/{first_id}/observations/batch", headers=headers,
                           json={"observed_at": observed_at, "entries": [
                               {"dish_id": dishes[0]["id"], "new_germinated_count": 0},
                               {"dish_id": dishes[1]["id"], "new_germinated_count": 2}]})
    assert observed.status_code == 200, observed.text
    assert {item["observed_at"] for item in observed.json()["created"]} == {observed_at}
    assert client.patch(f"/api/experiments/{first_id}/sowing/{dishes[0]['id']}", headers=headers,
                        json={"sown_at": "2026-09-30T00:00:00Z"}).status_code == 422
    assert client.patch(f"/api/experiments/{first_id}/sowing/{dishes[0]['id']}", headers=headers,
                        json={"sown_at": "2026-09-27T00:00:00Z"}).status_code == 200
    second = config(client, headers, [lot["id"]], name="第二次实验", replicates=1)
    second_id = second["experiment"]["id"]
    second_summary = ready(client, headers, second_id)
    assert second_summary["dishes"][0]["field_number"] == "001"
    cancelled = client.post(f"/api/experiments/{second_id}/sowing/{second_summary['dishes'][0]['id']}/cancel",
                            headers=headers, json={"reason": "培养皿破损"})
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["dishes"][0]["field_number"] == "001"
    assert cancelled.json()["cancelled_count"] == 1
    assert client.post(f"/api/experiments/{second_id}/observations/batch", headers=headers,
                       json={"observed_at": observed_at, "entries": [
                           {"dish_id": second_summary["dishes"][0]["id"], "new_germinated_count": 0}]}).status_code == 409
    workbook_response = client.post("/api/export/experiments/workbook.xlsx", headers=headers,
                                    json={"experiment_ids": [first_id, second_id]})
    assert workbook_response.status_code == 200, workbook_response.text
    workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
    assert workbook.sheetnames == ["01_材料总表", "02_发芽率汇总", "03_发芽原始记录",
                                   "04_幼苗测定长表", "05_幼苗测定宽表", "06_导出说明"]
    material_rows = list(workbook.worksheets[0].values)
    assert [item[0] for item in material_rows[1:]] == ["001", "002"]
    assert [item[2] for item in material_rows[1:]] == ["001", "001"]
    rate_rows = list(workbook.worksheets[1].values)
    assert sorted((row[7], row[8], row[9], row[10]) for row in rate_rows[1:]) == [(0, 0, 0, None), (2, 40, 2, 5)]
    observation_rows = list(workbook.worksheets[2].values)
    assert len(observation_rows) == 3
    assert sorted(item[9] for item in observation_rows[1:]) == [0, 2]
    assert observation_rows[1][3].startswith("001-")
    assert len(list(workbook.worksheets[3].values)) == 1
    assert list(workbook.worksheets[5].values)[1][0] == "导出时间"
    workbook.close()
    assert client.get(f"/api/experiments/{first_id}/configuration").json()["materials"][0]["experiment_number"] == 1
    single_response = client.post("/api/export/experiments/workbook.xlsx", headers=headers,
                                  json={"experiment_ids": [first_id]})
    assert single_response.status_code == 200
    single = load_workbook(BytesIO(single_response.content), read_only=True)
    assert len(list(single.worksheets[0].values)) == 2
    assert [item[0] for item in list(single.worksheets[0].values)[1:]] == ["001"]
    single.close()


def test_rate_summary_zero_germination_excludes_unplaced_and_cancelled_dishes(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers, json={"scientific_name": "Setaria viridis"}).json()
    lot = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxon["id"]}).json()
    experiment_id = config(client, headers, [lot["id"]], replicates=3)["experiment"]["id"]
    dishes = ready(client, headers, experiment_id)["dishes"]
    assert client.post(f"/api/experiments/{experiment_id}/sowing/batch", headers=headers,
                       json={"sown_at": "2026-09-28T09:00:00Z", "dish_ids": [dishes[0]["id"]]}).status_code == 200
    assert client.post(f"/api/experiments/{experiment_id}/sowing/{dishes[1]['id']}/cancel", headers=headers,
                       json={"reason": "培养皿破损"}).status_code == 200
    assert client.post(f"/api/experiments/{experiment_id}/observations/batch", headers=headers,
                       json={"observed_at": "2026-09-29T09:00:00Z", "entries": [
                           {"dish_id": dishes[0]["id"], "new_germinated_count": 0}]}).status_code == 200
    response = client.post("/api/export/experiments/workbook.xlsx", headers=headers,
                           json={"experiment_ids": [experiment_id]})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rate = list(workbook["02_发芽率汇总"].values)[1]
    assert rate[7:] == (1, 20, 0, 0)
    workbook.close()


def test_partial_daily_observation_keeps_other_dishes_pending(auth_client):
    client, headers = auth_client
    taxon = client.post("/api/taxa", headers=headers,
                        json={"scientific_name": "Setaria viridis", "common_name": "狗尾草"}).json()
    lot = client.post("/api/seed-lots", headers=headers, json={"taxon_id": taxon["id"]}).json()
    created = config(client, headers, [lot["id"]], name="批量巡检", replicates=100)
    experiment_id = created["experiment"]["id"]
    dishes = ready(client, headers, experiment_id)["dishes"]
    sow_at = datetime.now(timezone.utc) - timedelta(hours=3)
    response = client.post(f"/api/experiments/{experiment_id}/sowing/batch", headers=headers,
                           json={"sown_at": sow_at.isoformat(), "dish_ids": [dish["id"] for dish in dishes]})
    assert response.status_code == 200, response.text
    assert response.json()["today_pending_count"] == 100
    observed_at = (datetime.now(timezone.utc) - timedelta(hours=2)).isoformat()
    entries = [{"dish_id": dish["id"], "new_germinated_count": 0 if index == 0 else 2}
               for index, dish in enumerate(dishes[:26])]
    saved = client.post(f"/api/experiments/{experiment_id}/observations/batch", headers=headers,
                        json={"observed_at": observed_at, "entries": entries})
    assert saved.status_code == 200, saved.text
    assert len(saved.json()["created"]) == 26
    summary = client.get(f"/api/experiments/{experiment_id}/execution").json()
    assert summary["today_observed_count"] == 26
    assert summary["today_pending_count"] == 74
    assert summary["dishes"][0]["observation_count"] == 1
    assert summary["dishes"][0]["cumulative_germinated"] == 0
    assert summary["dishes"][1]["sample_count"] == 2
    assert summary["dishes"][26]["observation_count"] == 0
    again = client.post(f"/api/experiments/{experiment_id}/observations/batch", headers=headers,
                        json={"observed_at": datetime.now(timezone.utc).isoformat(), "entries": [
                            {"dish_id": dishes[1]["id"], "new_germinated_count": 1}]})
    assert again.status_code == 200, again.text
    assert client.get(f"/api/experiments/{experiment_id}/execution").json()["dishes"][1]["cumulative_germinated"] == 3


def test_catalog_order_and_dashboard_active_experiment_ids(auth_client):
    client, headers = auth_client
    taxa = [client.post("/api/taxa", headers=headers, json={
        "scientific_name": sci, "common_name": common}).json() for common, sci in
        [("狗尾草", "Setaria viridis"), ("白三叶", "Trifolium repens")]]
    lots = [client.post("/api/seed-lots", headers=headers,
                        json={"taxon_id": taxon["id"]}).json() for taxon in taxa]
    assert [item["common_name"] for item in client.get("/api/taxa").json()] == ["白三叶", "狗尾草"]
    assert [item["taxon"]["common_name"] for item in client.get("/api/seed-lots").json()] == ["白三叶", "狗尾草"]
    assert client.get("/api/dashboard").json()["active_experiment_ids"] == []
    ids = []
    for index, lot in enumerate(lots):
        created = config(client, headers, [lot["id"]], name=f"并行实验{index}", replicates=1)
        experiment_id = created["experiment"]["id"]
        dish = ready(client, headers, experiment_id)["dishes"][0]
        response = client.post(f"/api/experiments/{experiment_id}/sowing/batch", headers=headers,
                               json={"sown_at": "2026-09-29T08:00:00Z", "dish_ids": [dish["id"]]})
        assert response.status_code == 200, response.text
        ids.append(experiment_id)
        assert set(client.get("/api/dashboard").json()["active_experiment_ids"]) == set(ids)
    active = client.get("/api/experiments", params={"status": "active"}).json()
    assert {item["id"] for item in active} == set(ids)


def test_workflow_migration_preserves_existing_records_and_constraints(tmp_path, monkeypatch):
    db_path = tmp_path / "existing-v031.db"
    monkeypatch.setenv("SEEDLAB_DATABASE_URL", f"sqlite:///{db_path.as_posix()}")
    get_settings.cache_clear()
    config = Config(str(Path(__file__).resolve().parents[1] / "alembic.ini"))
    config.set_main_option("script_location", str(Path(__file__).resolve().parents[1] / "alembic"))
    try:
        command.upgrade(config, "a9c41e32b7d6")
        engine = make_engine(f"sqlite:///{db_path.as_posix()}")
        with engine.begin() as connection:
            connection.exec_driver_sql("INSERT INTO taxa (id, created_at, code, scientific_name, common_name, is_active) VALUES ('taxon', '2026-01-01', 'SP-0001', 'Setaria viridis', '狗尾草', 1)")
            connection.exec_driver_sql("INSERT INTO seed_lots (id, created_at, code, taxon_id, is_active) VALUES ('lot', '2026-01-01', 'LOT-2026-001', 'taxon', 1)")
            connection.exec_driver_sql("INSERT INTO experiments (id, created_at, code, name, status) VALUES ('exp', '2026-01-01', 'EXP-2026-001', '旧实验', 'active')")
            connection.exec_driver_sql("INSERT INTO experiment_materials (id, created_at, experiment_id, seed_lot_id, display_order) VALUES ('material', '2026-01-01', 'exp', 'lot', 0)")
            connection.exec_driver_sql("INSERT INTO germination_dishes (id, created_at, material_id, code, replicate_no, label, seed_count, sown_at) VALUES ('dish', '2026-01-01', 'material', 'EXP-2026-001-M001-R01', 1, 'R1', 20, '2026-01-01')")
            connection.exec_driver_sql("INSERT INTO germination_observations (id, created_at, dish_id, observed_at, new_germinated_count) VALUES ('observation', '2026-01-01', 'dish', '2026-01-02', 0)")
            connection.exec_driver_sql("INSERT INTO import_jobs (id, created_at, filename, status, total_rows, successful_rows) VALUES ('job', '2026-01-01', 'old.xlsx', 'completed', 1, 1)")
        engine.dispose()

        command.upgrade(config, "head")
        command.check(config)
        engine = make_engine(f"sqlite:///{db_path.as_posix()}")
        with engine.connect() as connection:
            assert connection.exec_driver_sql("SELECT version_num FROM alembic_version").scalar() == "c6d91f28a405"
            assert connection.exec_driver_sql("PRAGMA foreign_key_check").all() == []
            assert connection.exec_driver_sql("SELECT new_germinated_count FROM germination_observations WHERE id='observation'").scalar() == 0
            assert connection.exec_driver_sql("SELECT experiment_number FROM experiment_materials WHERE id='material'").scalar() is None
            assert connection.exec_driver_sql("SELECT sown_at FROM germination_dishes WHERE id='dish'").scalar() is not None
            assert connection.exec_driver_sql("SELECT source_code FROM seed_lots WHERE id='lot'").scalar() is None
        engine.dispose()

        engine = make_engine(f"sqlite:///{db_path.as_posix()}")
        with engine.begin() as connection:
            connection.exec_driver_sql("UPDATE experiment_materials SET experiment_number=1 WHERE id='material'")
            connection.exec_driver_sql("INSERT INTO seed_lots (id, created_at, code, taxon_id, is_active) VALUES ('lot2', '2026-01-01', 'LOT-2026-002', 'taxon', 1)")
            connection.exec_driver_sql("INSERT INTO experiment_materials (id, created_at, experiment_id, seed_lot_id, display_order) VALUES ('material2', '2026-01-01', 'exp', 'lot2', 1)")
        with pytest.raises(IntegrityError), engine.begin() as connection:
            connection.exec_driver_sql("UPDATE experiment_materials SET experiment_number=1 WHERE id='material2'")
        with pytest.raises(IntegrityError), engine.begin() as connection:
            connection.exec_driver_sql("UPDATE experiment_materials SET experiment_number=-1 WHERE id='material2'")
        with pytest.raises(IntegrityError), engine.begin() as connection:
            connection.exec_driver_sql("UPDATE germination_dishes SET cancelled_at='2026-01-03' WHERE id='dish'")
        engine.dispose()

        command.downgrade(config, "a9c41e32b7d6")
        command.upgrade(config, "head")
        engine = make_engine(f"sqlite:///{db_path.as_posix()}")
        with engine.connect() as connection:
            assert connection.exec_driver_sql("PRAGMA foreign_key_check").all() == []
            assert connection.exec_driver_sql("SELECT COUNT(*) FROM germination_observations WHERE id='observation'").scalar() == 1
            assert connection.exec_driver_sql("SELECT COUNT(*) FROM import_jobs WHERE id='job'").scalar() == 1
            assert connection.exec_driver_sql("PRAGMA index_list(import_jobs)").all()
        engine.dispose()

        fresh_path = tmp_path / "fresh.db"
        monkeypatch.setenv("SEEDLAB_DATABASE_URL", f"sqlite:///{fresh_path.as_posix()}")
        get_settings.cache_clear()
        command.upgrade(config, "head")
        fresh = make_engine(f"sqlite:///{fresh_path.as_posix()}")
        with fresh.connect() as connection:
            assert connection.exec_driver_sql("SELECT version_num FROM alembic_version").scalar() == "c6d91f28a405"
            assert connection.exec_driver_sql("PRAGMA foreign_key_check").all() == []
        fresh.dispose()
    finally:
        get_settings.cache_clear()
