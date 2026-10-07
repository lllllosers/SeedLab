"""Characterize existing lifecycle entry points and their audit snapshots."""
from datetime import datetime, timedelta, timezone
from io import BytesIO
from contextlib import contextmanager

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.db.session import get_db
from app.models import AuditLog

from test_experiment_config import design, make_lot
from test_seedling_measurement import observe, payload, setup_experiment


def audit_rows(client, entity_type, entity_id=None):
    # Use the existing isolated client session to inspect technical audit identity.
    # The public history intentionally exposes business labels rather than entity_id/user_id.
    query = select(AuditLog).where(AuditLog.entity_type == entity_type)
    if entity_id is not None:
        query = query.where(AuditLog.entity_id == entity_id)
    with contextmanager(client.app.dependency_overrides[get_db])() as db:
        stored = [{'id': row.id, 'action': row.action, 'user_id': row.user_id,
                   'before': row.before, 'after': row.after}
                  for row in db.scalars(query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc()))]
    response = client.get('/api/audit-logs', params={'page_size': 100, 'entity_type': entity_type})
    assert response.status_code == 200
    assert {row['id'] for row in stored} <= {row['id'] for row in response.json()['items']}
    return stored


def export_facts(client, headers, experiment_id):
    response = client.post('/api/export/experiments/workbook.xlsx', headers=headers,
                           json={'experiment_ids': [experiment_id]})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content), data_only=True)
    try:
        return {name: list(book[name].values) for name in book.sheetnames[:-1]}
    finally:
        book.close()


def test_configuration_numbering_and_sowing_keep_lifecycle_audit_snapshots(auth_client):
    client, headers = auth_client
    _, lot = make_lot(client, headers)
    data = design(lot['id'], dag_days=[0])
    data['description'] = None
    data['protocol'].update(replicate_count=2, sample_count=2, summary=None)
    data['materials'][0].update(label=None, seeds_per_dish_override=None,
                               replicate_count_override=None, sample_count_override=None)
    response = client.post('/api/experiments/configured', json=data, headers=headers)
    assert response.status_code == 201
    created = response.json()
    experiment = created['experiment']
    base = '/api/experiments/' + experiment['id']
    creation = audit_rows(client, 'Experiment', experiment['id'])[0]
    assert creation['action'] == 'create' and creation['before'] is None
    assert creation['after'] == {'code': experiment['code'], 'configuration': data}
    assert creation['user_id'] == experiment['owner_id']
    protocol = {**data['protocol'], 'sample_count': 3}
    assert client.put(base + '/protocol', json=protocol, headers=headers).status_code == 200
    protocol_audit = audit_rows(client, 'ExperimentProtocol')[0]
    assert protocol_audit['before']['sample_count'] == 2 and protocol_audit['after']['sample_count'] == 3
    material_id = created['materials'][0]['id']
    assert client.patch(base + '/materials/' + material_id,
                        json={'sample_count_override': 1}, headers=headers).status_code == 200
    material_audit = audit_rows(client, 'ExperimentMaterial', material_id)[0]
    assert material_audit['before']['sample_count_override'] is None
    assert material_audit['after']['sample_count_override'] == 1
    assert client.put(base + '/dag', json={'days': [2, 0]}, headers=headers).status_code == 200
    dag_audit = audit_rows(client, 'MeasurementTimepoint')[0]
    assert dag_audit['before'] == {'dag_days': [0]} and dag_audit['after'] == {'dag_days': [0, 2]}
    for status in ('ready', 'draft', 'ready'):
        result = client.patch(base, json={'status': status}, headers=headers)
        assert result.status_code == 200 and result.json()['status'] == status
        assert result.json()['started_at'] is result.json()['ended_at'] is None
    states = {(row['before']['status'], row['after']['status'])
              for row in audit_rows(client, 'Experiment', experiment['id']) if row['before']}
    assert {('draft', 'ready'), ('ready', 'draft')} <= states
    planned = client.post(base + '/confirm-numbers', headers=headers).json()
    assert planned['experiment']['status'] == 'ready' and planned['experiment']['numbering_locked_at']
    assert all(dish['sown_at'] is None for dish in planned['dishes'])
    locked = audit_rows(client, 'Experiment', experiment['id'])[0]
    assert locked['before'] == {'numbering_locked_at': None}
    assert locked['after']['numbering_locked_at'] and locked['after']['material_count'] == 1
    assert client.patch(base, json={'status': 'draft'}, headers=headers).status_code == 409
    assert client.post(base + '/reopen-design', headers=headers).status_code == 200
    unlocked = audit_rows(client, 'Experiment', experiment['id'])[0]
    assert unlocked['before'] == {'numbering_confirmed': True}
    assert unlocked['after'] == {'numbering_confirmed': False, 'status': 'draft'}
    assert client.get(base + '/execution').json()['dishes'] == []
    assert client.get(base + '/configuration').json()['materials'][0]['experiment_number'] is None
    assert client.patch(base, json={'status': 'ready'}, headers=headers).status_code == 200
    assert client.patch(base, json={'status': 'active'}, headers=headers).status_code == 409
    dishes = client.post(base + '/confirm-numbers', headers=headers).json()['dishes']
    for dish, when in ((dishes[1], '2026-09-03T08:00:00+08:00'), (dishes[0], '2026-09-01T08:00:00+08:00')):
        result = client.post(base + '/sowing/batch', headers=headers,
                             json={'dish_ids': [dish['id']], 'sown_at': when})
        assert result.status_code == 200 and result.json()['experiment']['status'] == 'active'
        sowing = audit_rows(client, 'GerminationDish', dish['id'])[0]
        assert sowing['before'] == {'sown_at': None}
        assert sowing['after']['sown_at'] == when[:10] + 'T00:00:00Z'
        assert audit_rows(client, 'Experiment', experiment['id'])[0]['after']['status'] == 'active'
    assert client.get(base + '/execution').json()['experiment']['started_at'] == '2026-09-01T00:00:00Z'
    assert client.patch(base + '/sowing/' + dishes[0]['id'], headers=headers,
                        json={'sown_at': '2026-09-02T08:00:00+08:00'}).status_code == 200
    corrected = audit_rows(client, 'GerminationDish', dishes[0]['id'])[0]
    assert corrected['before'] == {'sown_at': '2026-09-01T00:00:00Z'}
    assert corrected['after'] == {'sown_at': '2026-09-02T00:00:00Z'}
    start_audit = audit_rows(client, 'Experiment', experiment['id'])[0]
    assert start_audit['before'] == {'started_at': '2026-09-01T00:00:00Z'}
    assert start_audit['after'] == {'started_at': '2026-09-02T00:00:00Z'}
    audit_count = len(audit_rows(client, 'Experiment', experiment['id']))
    assert client.post(base + '/reopen-design', headers=headers).status_code == 409
    assert len(audit_rows(client, 'Experiment', experiment['id'])) == audit_count
    assert client.get(base).json()['code'] == experiment['code']
    assert [dish['field_number'] for dish in client.get(base + '/execution').json()['dishes']] == ['001-1', '001-2']


@pytest.mark.parametrize('ending', ['complete', 'patch_complete', 'terminate'])
def test_endings_preserve_facts_and_audit_then_apply_existing_correction_policy(auth_client, ending):
    client, headers = auth_client
    base, dish, _, _ = setup_experiment(client, headers, days=(0,))
    experiment_id = base.rsplit('/', 1)[-1]
    observe(client, headers, base, dish, datetime.now(timezone.utc) - timedelta(days=1))
    task = client.get(base + '/measurement-tasks').json()['tasks'][0]
    measurement = client.post(base + '/measurements', headers=headers,
                              json=payload(task, datetime.now(timezone.utc))).json()
    creation = audit_rows(client, 'SeedlingMeasurement', measurement['id'])[0]
    assert creation['before'] is None and creation['after']['root_length_mm'] == 0
    assert creation['after']['shoot_length_mm'] == 0
    facts_before = export_facts(client, headers, experiment_id)
    if ending == 'patch_complete':
        result = client.patch(base, json={'status': 'completed'}, headers=headers)
    elif ending == 'terminate':
        result = client.post(base + '/terminate', json={'reason': '  材料污染  '}, headers=headers)
    else:
        result = client.post(base + '/complete', headers=headers)
    assert result.status_code == 200
    ended = result.json()
    status = 'cancelled' if ending == 'terminate' else 'completed'
    assert ended['status'] == status and ended['ended_at'] and ended['started_at']
    assert ended['termination_reason'] == ('材料污染' if ending == 'terminate' else None)
    lifecycle = audit_rows(client, 'Experiment', experiment_id)[0]
    assert lifecycle['action'] == 'update' and lifecycle['before']['status'] == 'active'
    assert lifecycle['after']['status'] == status and lifecycle['after']['ended_at']
    if ending == 'terminate':
        assert lifecycle['after']['termination_reason'] == '材料污染'
    assert export_facts(client, headers, experiment_id) == facts_before
    assert client.get(base + '/measurement-tasks').json()['tasks'] == []
    assert client.get(base + '/measurement-records').json()['total'] == 1
    assert client.post(base + '/measurements', json=payload(task, datetime.now(timezone.utc)), headers=headers).status_code == 409
    assert client.delete(base + '/measurements/' + measurement['id'], headers=headers).status_code == 409
    corrected = client.patch(base + '/measurements/' + measurement['id'], json={'root_length_mm': 2}, headers=headers)
    if status == 'completed':
        assert corrected.status_code == 200
        change = audit_rows(client, 'SeedlingMeasurement', measurement['id'])[0]
        assert change['action'] == 'update' and change['before']['root_length_mm'] == 0
        assert change['after']['root_length_mm'] == 2 and change['after']['shoot_length_mm'] == 0
        assert client.patch(base + '/measurements/' + measurement['id'], json={'root_length_mm': 2}, headers=headers).status_code == 200
        assert len(audit_rows(client, 'SeedlingMeasurement', measurement['id'])) == 2
    else:
        assert corrected.status_code == 409
        assert len(audit_rows(client, 'SeedlingMeasurement', measurement['id'])) == 1
